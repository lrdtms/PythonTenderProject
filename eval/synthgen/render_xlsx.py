"""XLSX rendering with openpyxl. Records the sheet and cell range of each BoQ/quotation row."""

from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BOLD = Font(bold=True)
GREY = PatternFill("solid", fgColor="DDDDDD")


def write_boq(
    path: Path, title: str, rows: list[dict], messy: bool, rng: random.Random
) -> dict[str, tuple[str, str]]:
    """Write the pricing schedule. Returns item_number -> (sheet, cell_range)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BoQ"
    locations: dict[str, tuple[str, str]] = {}

    r = 1
    ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=13)
    if messy:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 1
        ws.cell(row=r, column=1, value="Synthetic test document - fictional data")
        r += 1
        ws.cell(
            row=r,
            column=1,
            value="Prepared by the employer's agent. Tenderers complete rate and amount columns.",
        )
        r += 2 + rng.randint(0, 2)
    else:
        r += 2
    header_row = r
    for c, h in enumerate(["Item", "Description", "Unit", "Qty", "Rate", "Amount"], start=1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = BOLD
        cell.fill = GREY
    r += 1

    last_section = None
    section_count = 0
    for row in rows:
        if row["section"] != last_section:
            if messy and last_section is not None:
                ws.cell(row=r, column=2, value=f"Subtotal: {last_section}").font = Font(italic=True)
                r += 1
            last_section = row["section"]
            section_count += 1
            ws.cell(row=r, column=1, value=f"SECTION: {row['section'].upper()}").font = BOLD
            if messy:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            r += 1
        ws.cell(row=r, column=1, value=row["item_number"])
        ws.cell(row=r, column=2, value=row["description"]).alignment = Alignment(wrap_text=True)
        ws.cell(row=r, column=3, value=row["raw_unit"])
        qty = row["raw_quantity"]
        if qty is not None:
            if messy and "," in qty:
                ws.cell(row=r, column=4, value=qty)  # text with a comma decimal
            else:
                try:
                    ws.cell(row=r, column=4, value=float(qty.replace(" ", "")))
                except ValueError:
                    ws.cell(row=r, column=4, value=qty)
        if row["unit_price_text"]:
            ws.cell(row=r, column=5, value=row["unit_price_text"])
        locations[row["item_number"]] = ("BoQ", f"A{r}:F{r}")
        r += 1

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 14
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if messy:
        notes = wb.create_sheet("Notes")
        notes["A1"] = "General notes"
        notes["A2"] = "Quantities are provisional and subject to remeasurement."
        notes["A3"] = "Item numbers refer to the specification in Part C3."
    wb.save(path)
    return locations


def write_quotation(
    path: Path,
    supplier: str,
    reference: str,
    date: str,
    valid_until: str,
    items: list[dict],
    currency: str = "ZAR",
) -> dict[int, tuple[str, str]]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Quote"
    ws["A1"] = supplier
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Quotation {reference}"
    ws["A3"] = f"Date: {date}"
    ws["A4"] = f"Valid until: {valid_until}"
    ws["A5"] = f"Currency: {currency}"
    hdr = 7
    for c, h in enumerate(["#", "Description", "Unit", "Qty", "Unit price"], start=1):
        cell = ws.cell(row=hdr, column=c, value=h)
        cell.font = BOLD
        cell.fill = GREY
    locations: dict[int, tuple[str, str]] = {}
    r = hdr + 1
    for it in items:
        ws.cell(row=r, column=1, value=it["row"])
        ws.cell(row=r, column=2, value=it["description"])
        ws.cell(row=r, column=3, value=it["unit"])
        if it["quantity"] is not None:
            ws.cell(row=r, column=4, value=float(it["quantity"]))
        if it["unit_price"] is not None:
            ws.cell(row=r, column=5, value=float(it["unit_price"]))
        locations[it["row"]] = ("Quote", f"A{r}:E{r}")
        r += 1
    ws.column_dimensions["B"].width = 80
    wb.save(path)
    return locations
