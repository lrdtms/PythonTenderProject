"""Check that the ground truth really matches the generated files."""

from __future__ import annotations

import json
import re
from pathlib import Path

import pymupdf
from openpyxl import load_workbook


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def check_pdf_quotes(pdf: Path, expected: list[tuple[str, str, int | None]]) -> list[str]:
    """expected: (id, quote, physical page). Returns failure messages."""
    failures = []
    with pymupdf.open(pdf) as doc:
        pages = [norm(p.get_text()) for p in doc]
    for gid, quote, page in expected:
        if page is None or not 1 <= page <= len(pages):
            failures.append(f"{gid}: page {page} out of range")
        elif norm(quote) not in pages[page - 1]:
            failures.append(f"{gid}: quote not found on page {page}")
    return failures


def verify_pack(pack_dir: Path) -> list[str]:
    gt = json.loads((pack_dir / "ground_truth.json").read_text(encoding="utf-8"))
    failures: list[str] = []
    docs = {d["filename"]: d for d in gt["documents"]}

    for name, d in docs.items():
        if not (pack_dir / name).exists():
            failures.append(f"missing file {name}")
            continue
        if d["format"] == "pdf" and d["scanned"]:
            with pymupdf.open(pack_dir / name) as pdf:
                if any(p.get_text().strip() for p in pdf):
                    failures.append(f"{name}: scanned PDF still has a text layer")
                if pdf.page_count != d["pages"]:
                    failures.append(f"{name}: page count changed")

    # Requirements and injection text live in the tender PDF.
    tender = next(d for d in gt["documents"] if d["kind"] == "tender_document")
    if not tender["scanned"]:
        expected = [(r["id"], r["evidence_quote"], r["source"]["page"]) for r in gt["requirements"]]
        expected += [(n["kind"], n["text"], n["source"]["page"]) for n in gt["non_requirements"]]
        failures += check_pdf_quotes(pack_dir / tender["filename"], expected)
    else:
        for r in gt["requirements"]:
            if not 1 <= r["source"]["page"] <= tender["pages"]:
                failures.append(f"{r['id']}: page out of range")

    # BoQ items
    for it in gt["boq_items"]:
        src = it["source"]
        path = pack_dir / src["document"]
        if src["sheet"]:
            ws = load_workbook(path, read_only=True)[src["sheet"]]
            row = next(iter(ws[src["cell_range"]]))
            if str(row[0].value) != it["item_number"] or norm(str(row[1].value)) != norm(
                it["description"]
            ):
                failures.append(f"BoQ {it['item_number']}: cell range does not match")
        elif not docs[src["document"]]["scanned"]:
            with pymupdf.open(path) as pdf:
                text = norm(pdf[src["page"] - 1].get_text())
            if it["item_number"].lower() not in text or norm(it["description"])[:30] not in text:
                failures.append(f"BoQ {it['item_number']}: not on page {src['page']}")

    # Quotations
    for q in gt["quotations"]:
        for qi in q["items"]:
            src = qi["source"]
            path = pack_dir / q["document"]
            if src["sheet"]:
                ws = load_workbook(path, read_only=True)[src["sheet"]]
                row = next(iter(ws[src["cell_range"]]))
                if norm(str(row[1].value)) != norm(qi["description"]):
                    failures.append(f"{q['document']} {src['cell_range']}: description mismatch")
            else:
                with pymupdf.open(path) as pdf:
                    text = norm(pdf[src["page"] - 1].get_text())
                if norm(qi["description"])[:30] not in text:
                    failures.append(f"{q['document']}: item not on page {src['page']}")
    return failures


def verify_all(root: Path) -> dict[str, list[str]]:
    return {
        p.name: verify_pack(p) for p in sorted(root.iterdir()) if (p / "ground_truth.json").exists()
    }
