"""F11 smoke: a 1,000-line estimate exports to Excel in under 10 s (renderer alone under 3 s), and
a 1,000-requirement checklist exports to PDF in under 10 s. There is no slow-test marker, so this
runs in every CI job; the budgets leave headroom and the measured figures are printed."""

from __future__ import annotations

import io
import random
import time
from datetime import UTC, datetime
from decimal import Decimal
from uuid import UUID, uuid4

from openpyxl import load_workbook

from app.db.session import session_scope
from app.kernel import OrgContext, Role
from app.modules.estimates.models import Estimate, EstimateLine
from app.modules.reporting.schemas import EstimateExportInput, ExportLine, ExportTotals
from app.modules.reporting.xlsx import LINES_SHEET, render_estimate_xlsx
from tests.conftest import make_pdf, make_tender, register, upload
from tests.test_exports_api import XLSX, _create, _download

LINES = 1000
SECTIONS = ("Preliminaries", "Cabling", "Lighting")
XLSX_BUDGET = 10.0
RENDER_BUDGET = 3.0
PDF_BUDGET = 10.0


def _seed_estimate(actor, tender_id: str, seed: int = 6) -> tuple[dict, list[dict]]:
    """An approved estimate with 1,000 lines through the ORM: random but valid snapshots in three
    sections, one line in ten excluded. Returns the estimate header values and the lines."""
    rng = random.Random(seed)  # noqa: S311 - deterministic test data
    user = UUID(actor.me["user_id"])
    org = UUID(actor.org_id)
    now = datetime.now(UTC)
    rows: list[dict] = []
    cost_sum = markup_sum = Decimal("0")
    for i in range(LINES):
        qty = Decimal(rng.randint(1, 500))
        price = Decimal(rng.randint(100, 900_000)) / 100
        row = {
            "position": i,
            "section": SECTIONS[i * len(SECTIONS) // LINES],
            "item_no": f"{i + 1}",
            "description": f"Line {i + 1} item {rng.randint(1000, 9999)}",
            "unit": rng.choice(["m", "ea", "lot"]),
            "quantity": qty,
        }
        if i % 10 == 9:
            row |= {"price_source": "excluded", "exclude_reason": "Client supplied"}
        else:
            cost = (qty * price).quantize(Decimal("0.01"))
            markup = (cost * Decimal("0.125")).quantize(Decimal("0.01"))
            cost_sum += cost
            markup_sum += markup
            row |= {
                "price_source": "manual",
                "unit_price_ex_vat": price.quantize(Decimal("0.0001")),
                "material_cost": cost,
                "labour_cost": Decimal("0.00"),
                "cost": cost,
                "markup": markup,
                "total": cost + markup,
            }
        rows.append(row)
    ex_vat = cost_sum + markup_sum
    vat = (ex_vat * Decimal("0.15")).quantize(Decimal("0.01"))
    header = {
        "subtotal_cost": cost_sum,
        "markup_total": markup_sum,
        "subtotal_ex_vat": ex_vat,
        "vat_total": vat,
        "total_inc_vat": ex_vat + vat,
        "at": now,
    }
    with session_scope() as s:
        est = Estimate(
            organisation_id=org, tender_id=UUID(tender_id), version=1, status="approved",
            currency="ZAR", vat_rate=Decimal("15.00"), markup_basis="cost",
            markup_percent=Decimal("12.50"), subtotal_cost=cost_sum, markup_total=markup_sum,
            subtotal_ex_vat=ex_vat, vat_total=vat, total_inc_vat=ex_vat + vat,
            created_by=user, reviewed_by=user, reviewed_at=now, approved_by=user,
            approved_at=now,
        )  # fmt: skip
        s.add(est)
        s.flush()
        s.add_all(
            EstimateLine(
                organisation_id=org, estimate_id=est.id, boq_item_id=uuid4(), updated_by=user, **r
            )
            for r in rows
        )
    return header, rows


def _export_input(header: dict, rows: list[dict]) -> EstimateExportInput:
    lines = tuple(
        ExportLine(
            position=r["position"], section=r["section"], item_no=r["item_no"],
            description=r["description"], unit=r["unit"], quantity=r["quantity"],
            price_source=r["price_source"], supplier_name=None, quoted_unit_price=None,
            quoted_currency=None, quoted_includes_vat=None,
            unit_price_ex_vat=r.get("unit_price_ex_vat"), labour_rate=None,
            material_cost=r.get("material_cost"), labour_cost=r.get("labour_cost"),
            cost=r.get("cost"),
            markup_percent=Decimal("12.50") if r["price_source"] == "manual" else None,
            markup=r.get("markup"), total=r.get("total"),
            status="excluded" if r["price_source"] == "excluded" else "priced",
            flags=(), exclude_reason=r.get("exclude_reason"),
        )
        for r in rows
    )  # fmt: skip
    return EstimateExportInput(
        tender_title="Perf tender", tender_reference="PERF-1", client_name=None, version=1,
        status="approved", currency="ZAR", vat_rate=Decimal("15.00"), markup_basis="cost",
        markup_percent=Decimal("12.50"),
        totals=ExportTotals(
            header["subtotal_cost"], header["markup_total"], header["subtotal_ex_vat"],
            header["vat_total"], header["total_inc_vat"],
        ),
        priced=LINES - LINES // 10, excluded=LINES // 10, unresolved=0, changed=0,
        incomplete=False, reviewed_by_name="Owner", reviewed_at=header["at"],
        approved_by_name="Owner", approved_at=header["at"], generated_by_name="Owner",
        lines=lines, generated_at=header["at"], export_id=uuid4(),
    )  # fmt: skip


def test_f11_estimate_workbook_of_1000_lines_is_fast() -> None:
    actor = register("perf1@example.com")
    tender_id = make_tender(actor)
    header, rows = _seed_estimate(actor, tender_id)

    started = time.perf_counter()
    rendered = render_estimate_xlsx(_export_input(header, rows))
    render_s = time.perf_counter() - started

    started = time.perf_counter()
    r = _create(actor, tender_id)
    post_s = time.perf_counter() - started
    print(f"F11 estimate_xlsx: render={render_s:.2f}s post={post_s:.2f}s ({LINES} lines)")
    assert r.status_code == 201, r.text
    assert r.json()["label"] == "approved" and r.json()["row_count"] == LINES
    assert post_s < XLSX_BUDGET
    assert render_s < RENDER_BUDGET

    download = _download(actor, r.json()["id"])
    assert download.status_code == 200 and download.headers["content-type"] == XLSX
    for blob in (download.content, rendered):
        ws = load_workbook(io.BytesIO(blob))[LINES_SHEET]
        assert ws.max_row == LINES + 1  # the header plus 1,000 lines; approved has no banner


def test_f11_checklist_pdf_of_1000_requirements_is_fast() -> None:
    from app.modules.checklist import public as checklist
    from app.modules.pipeline.models import ExtractionRun

    actor = register("perf2@example.com")
    tender_id = make_tender(actor)
    body = upload(actor, tender_id, make_pdf(["one"]), "tender.pdf").json()
    ctx = OrgContext(UUID(actor.me["user_id"]), UUID(actor.org_id), Role.ADMIN)
    rows = [
        {
            "description": f"The tenderer shall provide item {i}",
            "requirement_type": "document",
            "mandatory_status": "mandatory" if i % 2 else "conditional",
            "evidence_quote": f"quote {i}",
            "source_order": i,
        }
        for i in range(LINES)
    ]
    with session_scope() as s:
        run = ExtractionRun(
            organisation_id=ctx.organisation_id,
            document_version_id=UUID(body["version"]["id"]),
            job_id=UUID(str(body["job_id"])),
            kind="requirements",
        )
        s.add(run)
        s.flush()
        checklist.replace_unreviewed(s, ctx, run.document_version_id, run.id, rows)

    started = time.perf_counter()
    r = _create(actor, tender_id, kind="checklist_pdf")
    post_s = time.perf_counter() - started
    print(f"F11 checklist_pdf: post={post_s:.2f}s ({LINES} requirements)")
    assert r.status_code == 201, r.text
    assert r.json()["row_count"] == LINES
    assert post_s < PDF_BUDGET
