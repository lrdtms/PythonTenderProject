"""Excel estimate renderer: golden file, draft marking, injection, precision, boundaries."""

import ast
import io
import json
import re
import zipfile
from dataclasses import replace
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from uuid import UUID
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook

from app.modules.reporting.schemas import (
    EstimateExportInput,
    ExportLine,
    ExportTotals,
)
from app.modules.reporting.xlsx import (
    LINE_COLUMNS,
    LINES_SHEET,
    SUMMARY_SHEET,
    is_draft,
    render_estimate_xlsx,
)

FIXTURES = Path(__file__).parent / "fixtures" / "exports"
GOLDEN = json.loads((FIXTURES / "golden_estimate.json").read_text(encoding="utf-8"))
CORPUS = json.loads((FIXTURES / "injection_corpus.json").read_text(encoding="utf-8"))
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REPORTING = Path(__file__).parent.parent / "app" / "modules" / "reporting"
BANNER = "DRAFT: version 2 is not approved. Do not submit."


def _dec(v):
    return None if v is None else Decimal(v)


def _dt(v):
    return None if v is None else datetime.fromisoformat(v.replace("Z", "+00:00"))


def build_input(raw: dict) -> EstimateExportInput:
    lines = tuple(
        ExportLine(
            **{
                **{k: v for k, v in ln.items()},
                **{
                    k: _dec(ln[k])
                    for k in (
                        "quantity",
                        "quoted_unit_price",
                        "unit_price_ex_vat",
                        "labour_rate",
                        "material_cost",
                        "labour_cost",
                        "cost",
                        "markup_percent",
                        "markup",
                        "total",
                    )
                },
                "flags": tuple(ln["flags"]),
            }
        )
        for ln in raw["lines"]
    )
    totals = raw["totals"]
    return EstimateExportInput(
        **{
            **raw,
            "vat_rate": Decimal(raw["vat_rate"]),
            "markup_percent": Decimal(raw["markup_percent"]),
            "totals": None
            if totals is None
            else ExportTotals(**{k: Decimal(v) for k, v in totals.items()}),
            "reviewed_at": _dt(raw["reviewed_at"]),
            "approved_at": _dt(raw["approved_at"]),
            "generated_at": _dt(raw["generated_at"]),
            "export_id": UUID(raw["export_id"]),
            "lines": lines,
        }
    )


GOLDEN_INPUT = build_input(GOLDEN["input"])


def approved(data: EstimateExportInput = GOLDEN_INPUT) -> EstimateExportInput:
    return replace(
        data,
        status="approved",
        approved_by_name="Lindiwe Dlamini",
        approved_at=datetime.fromisoformat("2026-10-23T10:15:00+02:00"),
    )


def raw_cells(xlsx: bytes, sheet_no: int) -> dict[str, tuple[str | None, str]]:
    with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
        root = ET.fromstring(z.read(f"xl/worksheets/sheet{sheet_no}.xml"))  # noqa: S314
    out = {}
    for c in root.iter(f"{{{NS['m']}}}c"):
        t = c.get("t")
        if t == "inlineStr":
            text = "".join(x.text or "" for x in c.iter(f"{{{NS['m']}}}t"))
            out[c.get("r")] = (t, text)
        else:
            v = c.find("m:v", NS)
            if v is not None:
                out[c.get("r")] = (t, v.text or "")
    return out


def sheet_xml(xlsx: bytes, sheet_no: int) -> str:
    with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
        return z.read(f"xl/worksheets/sheet{sheet_no}.xml").decode("utf-8")


def cell_map(xlsx: bytes) -> dict[str, dict[str, tuple]]:
    wb = load_workbook(io.BytesIO(xlsx))
    return {
        ws.title: {
            c.coordinate: (c.value, c.number_format, c.data_type)
            for row in ws.iter_rows()
            for c in row
            if c.value is not None
        }
        for ws in wb.worksheets
    }


def display(raw_v: str) -> Decimal:
    return Decimal(format(float(raw_v), ".15g"))


def find_row(ws, text: str, col: int = 1) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, col).value == text:
            return row
    raise AssertionError(f"no row with {text!r} in column {col}")


@pytest.fixture(scope="module")
def golden_bytes() -> bytes:
    return render_estimate_xlsx(GOLDEN_INPUT)


def test_constants():
    assert (SUMMARY_SHEET, LINES_SHEET) == ("Summary", "Lines")
    assert len(LINE_COLUMNS) == 21
    assert LINE_COLUMNS[0] == "Line"
    assert LINE_COLUMNS[18] == "Line total (ex VAT)"
    assert LINE_COLUMNS[-1] == "Notes"
    assert is_draft("draft") and is_draft("reviewed") and not is_draft("approved")


# ---- golden ----


@pytest.mark.parametrize("sheet_no,name", [(1, "Summary"), (2, "Lines")])
def test_golden_cells(golden_bytes, sheet_no, name):
    wb = load_workbook(io.BytesIO(golden_bytes))
    assert wb.sheetnames == ["Summary", "Lines"]
    ws = wb[name]
    raw = raw_cells(golden_bytes, sheet_no)
    expected = GOLDEN["expected"][name]
    actual = {c.coordinate: c for row in ws.iter_rows() for c in row if c.value is not None}
    assert set(actual) == set(expected), (
        f"extra: {sorted(set(actual) - set(expected))} missing: {sorted(set(expected) - set(actual))}"
    )
    for addr, exp in expected.items():
        t, v = raw[addr]
        assert actual[addr].number_format == exp["number_format"], addr
        if exp["type"] == "n":
            assert t in (None, "n"), f"{addr} should be numeric"
            assert display(v) == Decimal(exp["value"]), addr
        else:
            assert t == "inlineStr", f"{addr} should be an inline string"
            assert v == exp["value"], addr


# ---- no formulas ----


def test_no_formulas_or_links(golden_bytes):
    with zipfile.ZipFile(io.BytesIO(golden_bytes)) as z:
        names = z.namelist()
        for n in names:
            if n.startswith("xl/worksheets/") and n.endswith(".xml"):
                assert "<f>" not in z.read(n).decode() and "<f " not in z.read(n).decode()
        assert not any("hyperlink" in n.lower() or "externalLink" in n for n in names)
        for n in names:
            if n.endswith(".xml"):
                assert b"<hyperlink" not in z.read(n)


# ---- totals tie ----


def test_summary_totals_equal_input_and_lines_sum(golden_bytes):
    wb = load_workbook(io.BytesIO(golden_bytes))
    summary, lines = wb["Summary"], wb["Lines"]
    raw = raw_cells(golden_bytes, 1)
    t = GOLDEN_INPUT.totals
    for label, expected in [
        ("Subtotal cost", t.subtotal_cost),
        ("Markup", t.markup_total),
        ("Subtotal ex VAT", t.subtotal_ex_vat),
        ("VAT", t.vat_total),
        ("Total incl VAT", t.total_inc_vat),
    ]:
        row = find_row(summary, label)
        assert display(raw[f"B{row}"][1]) == expected
    header = find_row(lines, "Line total (ex VAT)", 19)
    status_col, total_col = 20, 19
    total = sum(
        (
            Decimal(format(float(lines.cell(r, total_col).value), ".15g"))
            for r in range(header + 1, lines.max_row + 1)
            if lines.cell(r, status_col).value == "Priced"
        ),
        Decimal(0),
    )
    assert total == t.subtotal_ex_vat


# ---- excluded and unresolved lines ----


def test_excluded_and_unresolved_rows(golden_bytes):
    ws = load_workbook(io.BytesIO(golden_bytes))["Lines"]
    ex = find_row(ws, "Excluded", 20)
    assert ws.cell(ex, 21).value == "Supplied by the client"
    assert ws.cell(ex, 1).fill.fgColor.rgb.endswith("D9D9D9")
    assert all(ws.cell(ex, c).value is None for c in range(9, 20))
    un = find_row(ws, "Not priced", 20)
    assert ws.cell(un, 21).value == "It is not known whether the quoted price includes VAT."
    assert all(ws.cell(un, c).value is None for c in range(12, 20))
    ch = find_row(ws, "Source changed", 20)
    assert ws.cell(ch, 21).value == "Source changed since this line was priced."
    assert all(ws.cell(ch, c).value is None for c in range(12, 20))


# ---- draft label ----


def _marked(wb) -> None:
    for ws in wb.worksheets:
        assert ws["A1"].value == BANNER
        assert any(str(r).startswith("A1:") for r in ws.merged_cells.ranges)
        assert ws["A1"].fill.fgColor.rgb.endswith("C00000")
        assert ws.sheet_properties.tabColor.rgb.endswith("C00000")
        assert "DRAFT" in ws.oddHeader.center.text
    assert "DRAFT" in wb.properties.title


def _unmarked(wb) -> None:
    for ws in wb.worksheets:
        assert not (isinstance(ws["A1"].value, str) and ws["A1"].value.startswith("DRAFT"))
        assert ws.sheet_properties.tabColor is None
        assert not ws.oddHeader.center.text
    assert "DRAFT" not in wb.properties.title


@pytest.mark.parametrize("status", ["draft", "reviewed", "approved"])
def test_draft_label(status):
    data = approved() if status == "approved" else replace(GOLDEN_INPUT, status=status)
    wb = load_workbook(io.BytesIO(render_estimate_xlsx(data)))
    summary = wb["Summary"]
    row = find_row(summary, "Status")
    if status == "approved":
        _unmarked(wb)
        assert summary.cell(row, 2).value == "Approved"
        by = summary.cell(find_row(summary, "Approved by"), 2).value
        assert "Lindiwe Dlamini" in by and "2026-10-23 08:15 UTC" in by
        assert wb["Lines"]["A1"].value == "Line"
    else:
        _marked(wb)
        assert summary.cell(row, 2).value == (
            "Draft" if status == "draft" else "Reviewed: not yet approved"
        )
        assert wb["Lines"]["A2"].value == "Line"


# ---- incomplete ----


def test_incomplete_and_missing_totals():
    ws = load_workbook(io.BytesIO(render_estimate_xlsx(GOLDEN_INPUT)))["Summary"]
    row = find_row(ws, "Total incl VAT")
    assert ws.cell(row, 3).value == "Incomplete: excludes 2 lines"
    complete = replace(GOLDEN_INPUT, incomplete=False)
    ws = load_workbook(io.BytesIO(render_estimate_xlsx(complete)))["Summary"]
    assert ws.cell(row, 3).value is None
    ws = load_workbook(io.BytesIO(render_estimate_xlsx(replace(GOLDEN_INPUT, totals=None))))[
        "Summary"
    ]
    for label in ("Subtotal cost", "Markup", "Subtotal ex VAT", "VAT", "Total incl VAT"):
        assert ws.cell(find_row(ws, label), 2).value == "Not available"


# ---- injection ----


def _hostile(s: str) -> EstimateExportInput:
    line = replace(
        GOLDEN_INPUT.lines[0],
        section=s,
        item_no=s,
        description=s,
        unit=s,
        supplier_name=s,
    )
    excluded = replace(GOLDEN_INPUT.lines[5], exclude_reason=s, description=s)
    return replace(
        approved(),
        tender_title=s,
        tender_reference=s,
        client_name=s,
        reviewed_by_name=s,
        approved_by_name=s,
        generated_by_name=s,
        lines=(line, excluded),
    )


@pytest.mark.parametrize("s", CORPUS, ids=range(len(CORPUS)))
def test_injection_corpus(s):
    xlsx = render_estimate_xlsx(_hostile(s))
    for n in (1, 2):
        xml = sheet_xml(xlsx, n)
        assert "<f>" not in xml and "<f " not in xml
    wb = load_workbook(io.BytesIO(xlsx))
    seen = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f"
                if isinstance(c.value, str):
                    seen += 1
                    v = c.value
                    risky = v.lstrip().startswith(("=", "+", "-", "@")) or v.startswith(
                        ("\t", "\r")
                    )
                    assert v.startswith("'") or not risky, repr(v[:30])
                    assert len(v) <= 32_000
    assert seen > 10


# ---- precision ----


def _precision_input() -> EstimateExportInput:
    line = replace(
        GOLDEN_INPUT.lines[0],
        quantity=Decimal("8.174"),
        total=Decimal("91356.93"),
        material_cost=Decimal("9735177631648.39"),
        labour_rate=Decimal("10.5000"),
    )
    return replace(GOLDEN_INPUT, lines=(line,))


def test_precision_numeric_branch():
    xlsx = render_estimate_xlsx(_precision_input())
    raw = raw_cells(xlsx, 2)
    row = 3
    assert display(raw[f"F{row}"][1]) == Decimal("8.174")
    assert display(raw[f"S{row}"][1]) == Decimal("91356.93")
    assert display(raw[f"N{row}"][1]) == Decimal("9735177631648.39")
    assert display(raw[f"M{row}"][1]) == Decimal("10.5")
    assert raw[f"M{row}"][0] in (None, "n")
    summary = load_workbook(io.BytesIO(xlsx))["Summary"]
    assert not any(
        isinstance(r[0].value, str) and "more than 15 digits" in r[0].value
        for r in summary.iter_rows()
    )


def test_precision_text_branch():
    long_q = Decimal("12345678901234.5678")
    long_total = Decimal("12345678901234.56")
    line = replace(GOLDEN_INPUT.lines[0], quantity=long_q, total=long_total)
    totals = replace(GOLDEN_INPUT.totals, total_inc_vat=long_total)
    xlsx = render_estimate_xlsx(replace(GOLDEN_INPUT, lines=(line,), totals=totals))
    raw = raw_cells(xlsx, 2)
    assert raw["F3"] == ("inlineStr", "12345678901234.5678")
    assert raw["S3"] == ("inlineStr", "12345678901234.56")
    summary = load_workbook(io.BytesIO(xlsx))["Summary"]
    total_row = find_row(summary, "Total incl VAT")
    assert raw_cells(xlsx, 1)[f"B{total_row}"] == ("inlineStr", "12345678901234.56")
    notes = [r[0].value for r in summary.iter_rows() if isinstance(r[0].value, str)]
    assert "Some values have more than 15 digits and are shown as text." in notes


# ---- header and footer ----

TITLES = ["A & B", "&Z&F", '&"Arial,Bold"x', "T" * 300]


@pytest.mark.parametrize("title", TITLES)
def test_header_footer_escaping(title):
    xlsx = render_estimate_xlsx(replace(GOLDEN_INPUT, tender_title=title))
    escaped = title[:60].replace("&", "&&")
    for n in (1, 2):
        root = ET.fromstring(sheet_xml(xlsx, n))  # noqa: S314
        hf = root.find("m:headerFooter", NS)
        assert hf is not None
        footer = hf.find("m:oddFooter", NS).text
        header = hf.find("m:oddHeader", NS).text
        assert escaped in footer
        rest = footer.replace(escaped, "")
        assert not re.findall(r"&(?![LCRPN]| )", rest), rest
        for text in (footer, header):
            for section in re.split(r"&[LCR]", text):
                assert len(section) < 255
        assert "DRAFT" in header


# ---- datetimes ----


def test_aware_datetimes_as_text():
    data = replace(
        approved(),
        reviewed_at=datetime.fromisoformat("2026-10-22T11:30:00+02:00"),
        generated_at=datetime(2026, 10, 23, 12, 0, tzinfo=UTC),
    )
    xlsx = render_estimate_xlsx(data)
    texts = [v for _, v in raw_cells(xlsx, 1).values()]
    assert "2026-10-23 12:00 UTC" in texts
    assert any(v.endswith("2026-10-22 09:30 UTC") for v in texts)
    assert any(v.endswith("2026-10-23 08:15 UTC") for v in texts)
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.properties.created == datetime(2026, 10, 23, 12, 0)


# ---- float ban ----


def test_no_float_in_reporting_money_path():
    for name in ("xlsx.py", "schemas.py"):
        tree = ast.parse((REPORTING / name).read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            assert not (isinstance(node, ast.Name) and node.id == "float"), name
            assert not (isinstance(node, ast.Constant) and isinstance(node.value, float)), name


# ---- determinism ----


def test_render_is_deterministic():
    a, b = render_estimate_xlsx(GOLDEN_INPUT), render_estimate_xlsx(GOLDEN_INPUT)
    assert cell_map(a) == cell_map(b)
    wa, wb_ = load_workbook(io.BytesIO(a)), load_workbook(io.BytesIO(b))
    assert wa.properties.created == wb_.properties.created
    assert wa.properties.creator == "IntelTPS"


def test_layout_settings(golden_bytes):
    ws = load_workbook(io.BytesIO(golden_bytes))["Lines"]
    assert ws.freeze_panes == "A3"
    assert ws.auto_filter.ref == "A2:U10"
    assert ws.print_title_rows == "$2:$2"
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True


def test_requires_generated_at_and_export_id():
    with pytest.raises(ValueError):
        render_estimate_xlsx(replace(GOLDEN_INPUT, generated_at=None))
    with pytest.raises(ValueError):
        render_estimate_xlsx(replace(GOLDEN_INPUT, export_id=None))
