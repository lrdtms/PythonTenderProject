"""Reporting service: storage, row, audit, permissions, limits and the failure path (ADR-0007)."""

from __future__ import annotations

import hashlib
import io
from dataclasses import replace
from datetime import UTC
from decimal import Decimal
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from app.db.session import session_scope
from app.infrastructure.storage import LocalVolumeStorage
from app.kernel import Forbidden, NotFound, OrgContext, Role, ValidationFailed
from app.modules.audit.models import AuditEvent
from app.modules.estimates.models import Estimate
from app.modules.estimates.repo import EstimateRepo
from app.modules.reporting import public as reporting
from app.modules.reporting import service
from app.modules.reporting.repo import ExportRepo
from app.modules.reporting.schemas import ChecklistExportInput, ChecklistRow, ExportCreate
from tests.conftest import make_tender, register
from tests.test_reporting_xlsx import GOLDEN_INPUT, approved


def _ctx(actor, role: Role = Role.ADMIN) -> OrgContext:
    return OrgContext(UUID(actor.me["user_id"]), UUID(actor.org_id), role)


def _estimate(ctx: OrgContext, tender_id: UUID) -> Estimate:
    return Estimate(
        organisation_id=ctx.organisation_id,
        tender_id=tender_id,
        version=2,
        status="draft",
        currency="ZAR",
        vat_rate=Decimal("15.00"),
        markup_basis="cost",
        markup_percent=Decimal("10.00"),
        created_by=ctx.user_id,
    )


def _checklist(rows: int = 1, review: str = "confirmed") -> ChecklistExportInput:
    row = ChecklistRow(
        requirement_type="technical",
        mandatory_status="mandatory",
        description="Supply UPS units",
        review_status=review,
        document_title="Spec",
        page_number=3,
        evidence_verified=True,
    )
    return ChecklistExportInput(
        tender_title="Hall wiring",
        tender_reference=None,
        client_name=None,
        closing_at=None,
        generated_by_name="Owner",
        rows=(row,) * rows,
        rejected_count=0,
    )


@pytest.fixture
def storage(tmp_path: Path) -> LocalVolumeStorage:
    return LocalVolumeStorage(tmp_path / "files")


def _setup(email: str):
    actor = register(email)
    ctx, tender_id = _ctx(actor), UUID(make_tender(actor))
    with session_scope() as session:
        est_id = EstimateRepo(session, ctx).add(_estimate(ctx, tender_id)).id
    return actor, ctx, tender_id, est_id


def _make_xlsx(ctx, storage, tender_id, est_id, payload=GOLDEN_INPUT):
    with session_scope() as session:
        return service.create_export(
            session,
            ctx,
            storage,
            tender_id,
            "estimate_xlsx",
            payload,
            estimate_id=est_id,
            estimate_tender_id=tender_id,
        )


def _files(storage: LocalVolumeStorage) -> list[Path]:
    return [p for p in storage.root.rglob("*") if p.is_file()]


def test_estimate_export_row_file_and_label(storage) -> None:
    _, ctx, tender_id, est_id = _setup("svc1@example.com")
    dto = _make_xlsx(ctx, storage, tender_id, est_id)  # the golden fixture is "reviewed"
    assert dto.kind == "estimate_xlsx"
    assert dto.label == GOLDEN_INPUT.status == "reviewed"
    assert dto.estimate_id == est_id and dto.estimate_version == GOLDEN_INPUT.version
    assert dto.file_name.endswith(".xlsx") and "-DRAFT-" in dto.file_name
    assert dto.row_count == len(GOLDEN_INPUT.lines)
    assert dto.generated_by == ctx.user_id
    assert dto.generated_by_name == GOLDEN_INPUT.generated_by_name
    with session_scope() as session:
        row = ExportRepo(session, ctx).get(dto.id)
        stored = storage.open(row.file_key).read()
        assert (
            row.file_key == f"org/{ctx.organisation_id}/tenders/{tender_id}/exports/{dto.id}.xlsx"
        )
        assert row.sha256 == hashlib.sha256(stored).hexdigest()
        assert row.file_size == len(stored) == dto.file_size
        assert row.content_type == dto.content_type
        summary = load_workbook(io.BytesIO(stored))["Summary"]
        cells = {r[0].value: r[1].value for r in summary.iter_rows() if r[0].value}
        # the file prints the stored time (to the minute) and the export id
        assert cells["Generated at (UTC)"] == row.generated_at.astimezone(UTC).strftime(
            "%Y-%m-%d %H:%M UTC"
        )
        assert cells["Export reference"] == str(dto.id)


def test_approved_estimate_has_no_draft_marker(storage) -> None:
    _, ctx, tender_id, est_id = _setup("svc2@example.com")
    dto = _make_xlsx(ctx, storage, tender_id, est_id, approved())
    assert dto.label == "approved" and "DRAFT" not in dto.file_name


@pytest.mark.parametrize("status", ["draft", "reviewed"])
def test_unapproved_estimate_is_draft_named(storage, status) -> None:
    _, ctx, tender_id, est_id = _setup(f"svc3{status}@example.com")
    dto = _make_xlsx(ctx, storage, tender_id, est_id, replace(GOLDEN_INPUT, status=status))
    assert dto.label == status and "-DRAFT-" in dto.file_name


@pytest.mark.parametrize(
    ("review", "label", "draft"),
    [("confirmed", "approved", False), ("unreviewed", "draft", True)],
)
def test_checklist_export_label_follows_review(storage, review, label, draft) -> None:
    actor = register(f"svc4{label}@example.com")
    ctx, tender_id = _ctx(actor), UUID(make_tender(actor))
    with session_scope() as session:
        dto = service.create_export(
            session, ctx, storage, tender_id, "checklist_pdf", _checklist(review=review)
        )
    assert dto.label == label and dto.estimate_id is None and dto.estimate_version is None
    assert dto.file_name.endswith(".pdf") and ("-DRAFT-" in dto.file_name) is draft
    assert dto.content_type == "application/pdf"
    with session_scope() as session:
        row = ExportRepo(session, ctx).get(dto.id)
        assert storage.open(row.file_key).read().startswith(b"%PDF")


def test_audit_event_has_only_safe_keys(storage) -> None:
    _, ctx, tender_id, est_id = _setup("svc5@example.com")
    dto = _make_xlsx(ctx, storage, tender_id, est_id)
    with session_scope() as session:
        event = session.scalar(
            select(AuditEvent).where(
                AuditEvent.entity_id == dto.id, AuditEvent.action == "export_created"
            )
        )
        assert event.entity_type == "export" and event.user_id == ctx.user_id
        assert event.event_metadata == {
            "kind": "estimate_xlsx",
            "estimate_id": str(est_id),
            "version": GOLDEN_INPUT.version,
            "status": "reviewed",
            "count": len(GOLDEN_INPUT.lines),
        }


def test_checklist_audit_has_no_estimate(storage) -> None:
    actor = register("svc5b@example.com")
    ctx, tender_id = _ctx(actor), UUID(make_tender(actor))
    with session_scope() as session:
        dto = service.create_export(session, ctx, storage, tender_id, "checklist_pdf", _checklist())
    with session_scope() as session:
        event = session.scalar(select(AuditEvent).where(AuditEvent.entity_id == dto.id))
        assert event.event_metadata == {
            "kind": "checklist_pdf",
            "estimate_id": None,
            "version": None,
            "status": "approved",
            "count": 1,
        }


def test_reviewer_cannot_export_or_download(storage) -> None:
    actor, ctx, tender_id, est_id = _setup("svc6@example.com")
    dto = _make_xlsx(ctx, storage, tender_id, est_id)
    reviewer = _ctx(actor, Role.REVIEWER)
    with session_scope() as session:
        with pytest.raises(Forbidden):
            service.create_export(
                session, reviewer, storage, tender_id, "checklist_pdf", _checklist()
            )
        with pytest.raises(Forbidden):
            service.open_export_file(session, reviewer, storage, dto.id)
        assert service.get_export(session, reviewer, dto.id).id == dto.id  # READ is allowed


def test_payload_must_match_kind(storage) -> None:
    _, ctx, tender_id, est_id = _setup("svc7@example.com")
    with session_scope() as session:
        with pytest.raises(ValidationFailed):
            service.create_export(
                session,
                ctx,
                storage,
                tender_id,
                "estimate_xlsx",
                _checklist(),
                estimate_id=est_id,
                estimate_tender_id=tender_id,
            )
        with pytest.raises(ValidationFailed):
            service.create_export(session, ctx, storage, tender_id, "checklist_pdf", GOLDEN_INPUT)
    assert _files(storage) == []


def test_estimate_must_belong_to_the_tender(storage) -> None:
    _, ctx, tender_id, est_id = _setup("svc8@example.com")
    with session_scope() as session:
        with pytest.raises(ValidationFailed):
            service.create_export(
                session,
                ctx,
                storage,
                tender_id,
                "estimate_xlsx",
                GOLDEN_INPUT,
                estimate_id=est_id,
                estimate_tender_id=uuid4(),
            )
        with pytest.raises(ValidationFailed):  # an estimate export without an estimate id
            service.create_export(session, ctx, storage, tender_id, "estimate_xlsx", GOLDEN_INPUT)
    assert _files(storage) == []


def test_too_many_rows_is_refused_and_nothing_is_stored(storage) -> None:
    _, ctx, tender_id, est_id = _setup("svc9@example.com")
    assert service.MAX_ROWS == 5_000 and reporting.max_rows() == 5_000
    big = replace(GOLDEN_INPUT, lines=GOLDEN_INPUT.lines[:1] * 5_001)
    with session_scope() as session:
        with pytest.raises(ValidationFailed) as err:
            service.create_export(
                session,
                ctx,
                storage,
                tender_id,
                "estimate_xlsx",
                big,
                estimate_id=est_id,
                estimate_tender_id=tender_id,
            )
        assert err.value.code == "export_too_large"
        with pytest.raises(ValidationFailed) as err2:
            service.create_export(
                session, ctx, storage, tender_id, "checklist_pdf", _checklist(rows=5_001)
            )
        assert err2.value.code == "export_too_large"
    assert _files(storage) == []
    with session_scope() as session:
        assert list(ExportRepo(session, ctx).find()) == []


def test_cap_is_read_at_call_time(storage, monkeypatch) -> None:
    actor = register("svc9b@example.com")
    ctx, tender_id = _ctx(actor), UUID(make_tender(actor))
    monkeypatch.setattr("app.modules.reporting.service.MAX_ROWS", 2)
    assert reporting.max_rows() == 2
    with session_scope() as session, pytest.raises(ValidationFailed) as err:
        service.create_export(session, ctx, storage, tender_id, "checklist_pdf", _checklist(rows=3))
    assert err.value.code == "export_too_large"
    # rejected requirements are not rows, so they do not count toward the cap
    with session_scope() as session:
        dto = service.create_export(
            session,
            ctx,
            storage,
            tender_id,
            "checklist_pdf",
            replace(_checklist(rows=1), rejected_count=10_000),
        )
    assert dto.row_count == 1


def test_file_is_deleted_and_no_row_remains_when_the_row_insert_fails(storage) -> None:
    _, ctx, tender_id, _ = _setup("svc10@example.com")
    with pytest.raises(IntegrityError), session_scope() as session:
        # an estimate that does not exist: the composite foreign key refuses the row at flush,
        # after the file was stored
        service.create_export(
            session,
            ctx,
            storage,
            tender_id,
            "estimate_xlsx",
            GOLDEN_INPUT,
            estimate_id=uuid4(),
            estimate_tender_id=tender_id,
        )
    assert _files(storage) == []
    with session_scope() as session:
        assert list(ExportRepo(session, ctx).find()) == []
        events = select(AuditEvent).where(
            AuditEvent.organisation_id == ctx.organisation_id,
            AuditEvent.action == "export_created",
        )
        assert session.scalars(events).all() == []


def test_download_returns_exact_bytes_and_audits(storage) -> None:
    _, ctx, tender_id, est_id = _setup("svc11@example.com")
    dto = _make_xlsx(ctx, storage, tender_id, est_id)
    with session_scope() as session:
        got, fh = service.open_export_file(session, ctx, storage, dto.id)
        with fh:
            data = fh.read()
        assert got.id == dto.id
        key = ExportRepo(session, ctx).get(dto.id).file_key
        assert data == storage.open(key).read() and len(data) == dto.file_size
    with session_scope() as session:
        event = session.scalar(
            select(AuditEvent).where(
                AuditEvent.entity_id == dto.id, AuditEvent.action == "export_downloaded"
            )
        )
        assert event.event_metadata == {"kind": "estimate_xlsx"}


def test_missing_stored_file_is_not_found(storage) -> None:
    _, ctx, tender_id, est_id = _setup("svc12@example.com")
    dto = _make_xlsx(ctx, storage, tender_id, est_id)
    with session_scope() as session:
        storage.delete(ExportRepo(session, ctx).get(dto.id).file_key)
    with session_scope() as session, pytest.raises(NotFound):
        service.open_export_file(session, ctx, storage, dto.id)
    with session_scope() as session:
        downloads = select(AuditEvent).where(
            AuditEvent.entity_id == dto.id, AuditEvent.action == "export_downloaded"
        )
        assert session.scalars(downloads).all() == []


def test_checklist_export_refuses_an_estimate_tender(storage) -> None:
    actor = register("svc12b@example.com")
    ctx, tender_id = _ctx(actor), UUID(make_tender(actor))
    with session_scope() as session, pytest.raises(ValidationFailed):
        service.create_export(
            session,
            ctx,
            storage,
            tender_id,
            "checklist_pdf",
            _checklist(),
            estimate_tender_id=tender_id,
        )


def test_other_organisation_gets_not_found(storage) -> None:
    _, ctx, tender_id, est_id = _setup("svc13a@example.com")
    dto = _make_xlsx(ctx, storage, tender_id, est_id)
    other = _ctx(register("svc13b@example.com", org="Other"))
    with session_scope() as session:
        with pytest.raises(NotFound):
            service.get_export(session, other, dto.id)
        with pytest.raises(NotFound):
            service.open_export_file(session, other, storage, dto.id)


def test_second_export_makes_a_second_row_and_file(storage) -> None:
    _, ctx, tender_id, est_id = _setup("svc14@example.com")
    a = _make_xlsx(ctx, storage, tender_id, est_id)
    b = _make_xlsx(ctx, storage, tender_id, est_id)
    assert a.id != b.id and len(_files(storage)) == 2
    with session_scope() as session:
        assert len(ExportRepo(session, ctx).find()) == 2


def test_export_create_schema() -> None:
    est = uuid4()
    assert ExportCreate(kind="estimate_xlsx", estimate_id=est).estimate_id == est
    assert ExportCreate(kind="checklist_pdf").estimate_id is None
    with pytest.raises(ValueError):
        ExportCreate(kind="checklist_pdf", estimate_id=est)
    with pytest.raises(ValueError):
        ExportCreate(kind="checklist_pdf", other=1)  # type: ignore[call-arg]
    with pytest.raises(ValueError):
        ExportCreate(kind="docx")  # type: ignore[arg-type]
