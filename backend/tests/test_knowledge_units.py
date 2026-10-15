"""Unit tests for the knowledge parsers, OCR helpers and chunking. No database, no API."""

from __future__ import annotations

import pymupdf
import pytest

from app.config import get_settings
from app.kernel import ProcessingFailed
from app.modules.knowledge.chunking import MAX_CHARS, ChunkDraft, chunk_document
from app.modules.knowledge.ocr import OcrFailed, TesseractOcr, parse_tsv
from app.modules.knowledge.parsers import parse_document, parse_pdf
from tests.conftest import (
    make_csv,
    make_docx,
    make_encrypted_pdf,
    make_pdf,
    make_scanned_pdf,
    make_table_pdf,
    make_xlsx,
)


class FakeOcr:
    def __init__(
        self, text: str = "Scanned text", confidence: float = 90.0, *, ready: bool = True
    ) -> None:
        self.text = text
        self.confidence = confidence
        self.ready = ready
        self.calls = 0

    def available(self) -> bool:
        return self.ready

    def run(self, png: bytes) -> tuple[str, float]:
        self.calls += 1
        assert png.startswith(b"\x89PNG")
        return self.text, self.confidence


class FailingOcr(FakeOcr):
    def run(self, png: bytes) -> tuple[str, float]:
        raise OcrFailed("exit_1")


# --- PDF ---------------------------------------------------------------------------------------


def test_text_pdf_yields_one_page_per_pdf_page_and_reports_progress() -> None:
    data = make_pdf(
        ["First page about cables and conduits.", "Second page about switchgear and panels."]
    )
    calls: list[tuple[int, int]] = []

    pages = parse_pdf(data, None, progress=lambda done, total: calls.append((done, total)))

    assert [p.page_number for p in pages] == [1, 2]
    assert "First page about cables" in pages[0].text
    assert "Second page about switchgear" in pages[1].text
    assert all(p.text_source == "embedded" and p.status == "ok" for p in pages)
    assert calls == [(1, 2), (2, 2)]


def test_pdf_with_ruled_table_records_table_rows() -> None:
    rows = [["Item", "Qty"], ["Cable", "10"], ["Switch", "4"]]
    pages = parse_pdf(make_table_pdf(rows), None)

    assert len(pages) == 1
    assert pages[0].status == "ok"
    assert "Cable" in pages[0].text
    assert len(pages[0].tables) == 1
    table = pages[0].tables[0]
    assert table["kind"] == "table"
    assert table["rows"] == rows


def test_scanned_pdf_uses_ocr_text_and_confidence() -> None:
    ocr = FakeOcr("Scanned text", 90.0)

    (page,) = parse_pdf(make_scanned_pdf(["anything"]), ocr)

    assert ocr.calls == 1
    assert page.text == "Scanned text"
    assert page.text_source == "ocr"
    assert page.status == "ok"
    assert page.ocr_confidence == 90.0


def test_scanned_pdf_with_low_confidence_is_flagged() -> None:
    (page,) = parse_pdf(make_scanned_pdf(["anything"]), FakeOcr("Murky text", 40.0))

    assert page.status == "low_confidence"
    assert page.text_source == "ocr"
    assert page.text == "Murky text"
    assert page.ocr_confidence == 40.0


def test_scanned_pdf_when_ocr_fails_has_empty_text() -> None:
    (page,) = parse_pdf(make_scanned_pdf(["anything"]), FailingOcr())

    assert page.status == "ocr_failed"
    assert page.text == ""
    assert page.ocr_confidence is None


def test_scanned_pdf_without_usable_ocr_is_unavailable() -> None:
    data = make_scanned_pdf(["anything"])

    (no_runner,) = parse_pdf(data, None)
    (not_ready,) = parse_pdf(data, FakeOcr(ready=False))

    for page in (no_runner, not_ready):
        assert page.status == "ocr_unavailable"
        assert page.text_source == "none"
        assert page.text == ""


def test_blank_pdf_page_is_empty() -> None:
    doc = pymupdf.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()

    (page,) = parse_pdf(data, None)

    assert page.status == "empty"
    assert page.text_source == "none"
    assert page.text == ""


def test_encrypted_pdf_is_a_permanent_failure() -> None:
    with pytest.raises(ProcessingFailed) as caught:
        parse_pdf(make_encrypted_pdf(), None)

    assert caught.value.error_code == "PDF_ENCRYPTED"
    assert caught.value.transient is False


def test_garbage_pdf_is_corrupt() -> None:
    with pytest.raises(ProcessingFailed) as caught:
        parse_pdf(b"%PDF-1.4 this is not really a pdf", None)

    assert caught.value.error_code == "FILE_CORRUPT"
    assert caught.value.transient is False


# --- XLSX, CSV, DOCX ---------------------------------------------------------------------------


def test_xlsx_gives_one_page_per_sheet() -> None:
    import io

    from openpyxl import Workbook

    wb = Workbook()
    first = wb.active
    first.title = "BoQ"
    first.append(["Item", "Qty"])
    first.append(["Cable", 10])
    wb.create_sheet("Notes")  # left empty
    buf = io.BytesIO()
    wb.save(buf)

    pages = parse_document("xlsx", buf.getvalue())

    assert [p.page_number for p in pages] == [1, 2]
    assert [p.label for p in pages] == ["BoQ", "Notes"]
    assert pages[0].text == "Item | Qty\nCable | 10"
    assert pages[0].status == "ok"
    assert pages[0].tables[0]["rows"] == [["Item", "Qty"], ["Cable", "10"]]
    assert pages[0].tables[0]["first_row"] == 1
    assert pages[0].tables[0]["kind"] == "sheet"
    assert pages[1].status == "empty"
    assert pages[1].text == ""


def test_xlsx_helper_builds_single_named_sheet() -> None:
    (page,) = parse_document("xlsx", make_xlsx([["a", "b"]], sheet="Prices"))

    assert page.label == "Prices"
    assert page.text == "a | b"


def test_csv_with_semicolons_and_cp1252_parses_into_rows() -> None:
    data = make_csv([["caf\xe9", "qty"], ["tea", "5"], ["milk", "7"]])
    assert b"caf\xe9;qty" in data

    (page,) = parse_document("csv", data)

    assert page.page_number == 1
    assert page.status == "ok"
    assert page.tables[0]["rows"] == [["caf\xe9", "qty"], ["tea", "5"], ["milk", "7"]]
    assert page.text.splitlines()[0] == "caf\xe9 | qty"


def test_docx_captures_paragraphs_and_table_on_a_single_page() -> None:
    data = make_docx(
        ["Scope of works", "Supply and install cabling."], table=[["Item", "Qty"], ["Cable", "10"]]
    )

    pages = parse_document("docx", data)

    assert len(pages) == 1
    page = pages[0]
    assert page.page_number == 1
    assert page.text == "Scope of works\nSupply and install cabling."
    assert page.tables == [{"kind": "table", "rows": [["Item", "Qty"], ["Cable", "10"]]}]
    assert page.label is not None
    assert "no page numbers" in page.label


def test_unsupported_format_is_rejected() -> None:
    with pytest.raises(ProcessingFailed) as caught:
        parse_document("exe", b"")

    assert caught.value.error_code == "UNSUPPORTED_FORMAT"


# --- OCR ---------------------------------------------------------------------------------------

TSV_HEADER = (
    "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext"
)


def _tsv_row(level: int, line: int, conf: float, text: str) -> str:
    return f"{level}\t1\t1\t1\t{line}\t1\t0\t0\t10\t10\t{conf}\t{text}"


def test_parse_tsv_rebuilds_lines_and_mean_confidence() -> None:
    tsv = "\n".join(
        [
            TSV_HEADER,
            _tsv_row(4, 1, -1, ""),  # a line-level row is ignored
            _tsv_row(5, 1, 90, "Hello"),
            _tsv_row(5, 1, 80, "world"),
            _tsv_row(5, 2, 70, "Second"),
            _tsv_row(5, 2, -1, "ignored"),  # -1 confidence: text kept, score ignored
            _tsv_row(5, 2, 60, "line"),
        ]
    )

    text, confidence = parse_tsv(tsv)

    assert text == "Hello world\nSecond ignored line"
    assert confidence == pytest.approx((90 + 80 + 70 + 60) / 4)


def test_parse_tsv_empty_input() -> None:
    assert parse_tsv("") == ("", 0.0)
    assert parse_tsv(TSV_HEADER) == ("", 0.0)


def test_tesseract_with_missing_executable() -> None:
    ocr = TesseractOcr(command="definitely-not-a-real-tesseract-binary")

    assert ocr.available() is False
    with pytest.raises(OcrFailed):
        ocr.run(b"\x89PNG not really")


# --- chunking ----------------------------------------------------------------------------------


def _long_page(prefix: str = "Clause", count: int = 120) -> str:
    return " ".join(f"{prefix} {i} applies to the supplier." for i in range(count))


def test_chunks_never_cross_pages() -> None:
    pages = [(1, _long_page("Alpha")), (2, "A short second page."), (3, _long_page("Gamma"))]

    chunks = chunk_document(pages)

    assert {c.page_number for c in chunks} == {1, 2, 3}
    for chunk in chunks:
        if chunk.page_number == 1:
            assert "Alpha" in chunk.text and "Gamma" not in chunk.text
        elif chunk.page_number == 2:
            assert chunk.text == "A short second page."
        else:
            assert "Gamma" in chunk.text and "Alpha" not in chunk.text


def test_long_page_is_split_into_bounded_chunks() -> None:
    chunks = chunk_document([(1, _long_page())])

    assert len(chunks) > 1
    assert all(len(c.text) <= MAX_CHARS * 2 for c in chunks)
    # an unbroken run with no sentence breaks is still capped
    (blob,) = chunk_document([(1, "x" * (MAX_CHARS * 4))])
    assert len(blob.text) == MAX_CHARS * 2


def test_consecutive_chunks_overlap_by_one_sentence() -> None:
    chunks = chunk_document([(1, _long_page())])

    assert len(chunks) >= 2
    for before, after in zip(chunks, chunks[1:], strict=False):
        last_sentence = "Clause " + before.text.rsplit("Clause ", 1)[1]
        assert after.text.startswith(last_sentence)
        # ...but only one sentence is repeated: the one before it is not
        previous = "Clause " + before.text.rsplit("Clause ", 2)[1]
        assert previous not in after.text


def test_heading_becomes_section_and_carries_to_next_page() -> None:
    pages = [
        (1, "3.2 Eligibility\n\nBidders must be registered with the authority."),
        (2, "Bidders must also hold a valid tax certificate."),
        (3, "SECTION B: CABLES\n\nCables must be armoured."),
    ]

    chunks = chunk_document(pages)

    by_page = {c.page_number: c for c in chunks}
    assert by_page[1].section == "3.2 Eligibility"
    assert by_page[2].section == "3.2 Eligibility"
    assert by_page[3].section == "SECTION B: CABLES"


def test_text_before_any_heading_has_no_section() -> None:
    (chunk,) = chunk_document([(1, "Plain opening text with no heading.")])

    assert chunk.section is None


def test_empty_pages_produce_no_chunks() -> None:
    assert chunk_document([(1, ""), (2, "   \n\n "), (3, "Real content here.")]) == [
        ChunkDraft(3, None, "Real content here.")
    ]
    assert chunk_document([]) == []


def test_chunking_is_deterministic() -> None:
    pages = [(1, "1.1 Scope\n\n" + _long_page()), (2, _long_page("Beta"))]

    assert chunk_document(pages) == chunk_document(pages)


# --- resource caps (I-1) -----------------------------------------------------------------------


def _xlsx_with_dimension(rows: list[list[object]], ref: str) -> bytes:
    """A workbook whose sheet XML declares a forged <dimension ref=...>."""
    import io
    import re
    import zipfile

    source = zipfile.ZipFile(io.BytesIO(make_xlsx(rows)))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data, count = re.subn(
                    rb'<dimension ref="[^"]*"', f'<dimension ref="{ref}"'.encode(), data
                )
                assert count == 1
            target.writestr(item, data)
    return out.getvalue()


def test_xlsx_forged_dimension_does_not_pad_rows() -> None:
    data = _xlsx_with_dimension([["Item", "Qty"], ["Cable", 10], ["Switch", 4]], "A1:XFD3")

    (page,) = parse_document("xlsx", data)

    assert page.tables[0]["rows"] == [["Item", "Qty"], ["Cable", "10"], ["Switch", "4"]]
    assert page.text == "Item | Qty\nCable | 10\nSwitch | 4"


def test_xlsx_trailing_empty_cells_are_trimmed() -> None:
    (page,) = parse_document("xlsx", make_xlsx([["a", None, "b", None, None], ["c"]]))

    assert page.tables[0]["rows"] == [["a", "", "b"], ["c"]]


def test_xlsx_wider_than_the_column_cap_fails_permanently() -> None:
    assert get_settings().xlsx_max_columns == 200
    ok = parse_document("xlsx", make_xlsx([[f"c{i}" for i in range(200)]]))
    assert len(ok[0].tables[0]["rows"][0]) == 200

    with pytest.raises(ProcessingFailed) as caught:
        parse_document("xlsx", make_xlsx([[f"c{i}" for i in range(201)]]))

    assert caught.value.error_code == "SHEET_TOO_LARGE"
    assert caught.value.transient is False
    assert "c200" not in str(caught.value)


def test_xlsx_with_more_cells_than_the_cap_fails_permanently(monkeypatch) -> None:
    assert get_settings().xlsx_max_cells == 2_000_000
    monkeypatch.setattr(get_settings(), "xlsx_max_cells", 50)
    rows = [[f"r{i}", i] for i in range(30)]  # 60 cells

    with pytest.raises(ProcessingFailed) as caught:
        parse_document("xlsx", make_xlsx(rows))

    assert caught.value.error_code == "SHEET_TOO_LARGE"
    assert parse_document("xlsx", make_xlsx(rows[:25]))[0].status == "ok"  # 50 cells is allowed


def _image_page_pdf(width: float, height: float) -> bytes:
    """One scanned-looking page (a small image, no text) on a page of the given size in points."""
    tiny = pymupdf.open()
    tiny.new_page(width=50, height=50).insert_text((5, 25), "x")
    png = tiny[0].get_pixmap(dpi=20).tobytes("png")
    tiny.close()
    doc = pymupdf.open()
    page = doc.new_page(width=width, height=height)
    page.insert_image(pymupdf.Rect(0, 0, 200, 200), stream=png)
    data = doc.tobytes()
    doc.close()
    return data


class _PixmapSpy:
    """Records the DPI asked of get_pixmap without allocating a real image."""

    def __init__(self) -> None:
        self.dpis: list[int] = []

    def __call__(self, *args, **kwargs):  # set on the class: not bound, so no page argument
        self.dpis.append(kwargs.get("dpi"))
        return type("Pix", (), {"tobytes": staticmethod(lambda fmt="png": b"\x89PNG fake")})()


def test_huge_scanned_page_is_not_rendered(monkeypatch) -> None:
    data = _image_page_pdf(14400, 14400)  # built before get_pixmap is replaced
    spy = _PixmapSpy()
    monkeypatch.setattr(pymupdf.Page, "get_pixmap", spy)
    ocr = FakeOcr()

    (page,) = parse_pdf(data, ocr, ocr_dpi=300)

    assert spy.dpis == []  # 60,000 px square at 300 DPI; even 100 DPI is far above the cap
    assert ocr.calls == 0
    assert page.status == "ocr_failed" and page.text_source == "none"


def test_large_scanned_page_is_rendered_at_a_lower_dpi(monkeypatch) -> None:
    data = _image_page_pdf(3000, 3000)
    spy = _PixmapSpy()
    monkeypatch.setattr(pymupdf.Page, "get_pixmap", spy)

    (page,) = parse_pdf(data, FakeOcr(), ocr_dpi=300)

    (dpi,) = spy.dpis
    assert 100 <= dpi < 300
    pixels = (3000 / 72 * dpi) ** 2
    assert pixels <= get_settings().ocr_max_megapixels * 1_000_000
    assert page.status == "ok" and page.text_source == "ocr"


def test_normal_scanned_page_keeps_the_configured_dpi(monkeypatch) -> None:
    data = _image_page_pdf(595, 842)  # A4
    spy = _PixmapSpy()
    monkeypatch.setattr(pymupdf.Page, "get_pixmap", spy)

    parse_pdf(data, FakeOcr(), ocr_dpi=300)

    assert spy.dpis == [300]


def test_tesseract_runs_single_threaded(monkeypatch) -> None:
    import subprocess

    from app.modules.knowledge import ocr as ocr_module

    seen: dict = {}

    def fake_run(args, **kwargs):
        seen.update(kwargs)
        return subprocess.CompletedProcess(args, 0, stdout=TSV_HEADER.encode(), stderr=b"")

    monkeypatch.setenv("PATH_MARKER_FOR_TEST", "kept")
    monkeypatch.setattr(ocr_module.subprocess, "run", fake_run)

    TesseractOcr().run(b"\x89PNG")

    assert seen["env"]["OMP_THREAD_LIMIT"] == "1"
    assert seen["env"]["PATH_MARKER_FOR_TEST"] == "kept"  # built from os.environ
