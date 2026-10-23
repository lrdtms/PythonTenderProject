"""Export routes: roles, version rules, file content, frozen exports, races, headers, audit."""

from __future__ import annotations

import io
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from uuid import UUID

import pdfplumber
import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import session_scope
from app.infrastructure.storage import get_storage
from app.kernel import OrgContext, Role
from app.modules.audit.models import AuditEvent
from app.modules.estimates import public as estimates
from app.modules.reporting.models import Export
from app.modules.suppliers import public as suppliers
from tests.conftest import add_member, make_pdf, make_tender, register, upload
from tests.test_estimates_api import (
    API,
    TENDERS,
    _lines,
    _move,
    _priced_setup,
    _ready,
    _resolve_all,
    _set_quantity,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORTS = f"{API}/exports"


def _create(actor, tender_id: str, **body):
    body.setdefault("kind", "estimate_xlsx")
    return actor.client.post(f"{TENDERS}/{tender_id}/exports", json=body)


def _download(actor, export_id: str):
    return actor.client.get(f"{EXPORTS}/{export_id}/download")


def _book(response):
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content))


def _approved(email: str):
    actor, tender_id, est = _ready(email)
    assert _move(actor, est, "review").status_code == 200
    assert _move(actor, est, "approve").status_code == 200
    return actor, tender_id, est


def _summary(wb) -> dict[str, object]:
    return {r[0].value: r[1].value for r in wb["Summary"].iter_rows() if r[0].value}


def _sheet_values(ws) -> list[tuple]:
    return [tuple(c.value for c in row) for row in ws.iter_rows()]


def _data_rows(ws) -> list[tuple]:
    """Lines rows below the header (a draft sheet has a banner row above it)."""
    rows = _sheet_values(ws)
    start = next(i for i, r in enumerate(rows) if r[0] == "Line") + 1
    return rows[start:]


def _dec(value) -> Decimal:
    return Decimal(str(value))


# --- roles -----------------------------------------------------------------------------------


def test_roles_for_create_download_and_read() -> None:
    admin, tender_id, _ = _approved("xr1@example.com")
    estimator = add_member(admin, "xr1e@example.com", "estimator")
    manager = add_member(admin, "xr1m@example.com", "project_manager")
    reviewer = add_member(admin, "xr1r@example.com", "reviewer")
    for who in (admin, estimator, manager):
        r = _create(who, tender_id)
        assert r.status_code == 201, r.text
        assert r.json()["generated_by_name"]
    made = _create(admin, tender_id).json()
    assert _create(reviewer, tender_id).status_code == 403
    assert _download(reviewer, made["id"]).status_code == 403
    got = reviewer.client.get(f"{EXPORTS}/{made['id']}")
    assert got.status_code == 200
    assert got.json()["generated_by_name"] == "Owner"
    assert _download(manager, made["id"]).status_code == 200


# --- version rules (D2) ----------------------------------------------------------------------


def test_default_version_rules() -> None:
    actor = register("xv1@example.com")
    tender_id = make_tender(actor)
    none = _create(actor, tender_id)
    assert none.status_code == 409 and none.json()["error"]["code"] == "no_estimate"

    actor, tender_id, est = _priced_setup("xv2@example.com")
    draft = _create(actor, tender_id)
    assert draft.status_code == 201, draft.text
    assert draft.json()["label"] == "draft" and draft.json()["estimate_version"] == 1

    _resolve_all(actor, est)
    assert _move(actor, est, "review").status_code == 200
    assert _move(actor, est, "approve").status_code == 200
    v2 = actor.client.post(f"{TENDERS}/{tender_id}/estimates").json()  # a newer draft
    assert v2["version"] == 2
    default = _create(actor, tender_id).json()
    assert default["label"] == "approved" and default["estimate_version"] == 1
    explicit = _create(actor, tender_id, estimate_id=v2["id"]).json()
    assert explicit["label"] == "draft" and explicit["estimate_version"] == 2
    assert "DRAFT" in explicit["file_name"] and "DRAFT" not in default["file_name"]


def test_estimate_id_must_belong_to_the_tender_and_kind() -> None:
    a, tender_a, est_a = _priced_setup("xv3a@example.com")
    other = make_tender(a, "Second")
    assert _create(a, other, estimate_id=est_a["id"]).status_code == 404
    b = register("xv3b@example.com")
    tender_b = make_tender(b)
    assert _create(b, tender_b, estimate_id=est_a["id"]).status_code == 404
    assert _create(b, tender_a).status_code == 404
    assert _create(a, tender_a, kind="checklist_pdf", estimate_id=est_a["id"]).status_code == 422
    assert _create(a, tender_a, kind="bogus").status_code == 422
    assert a.client.post(f"{TENDERS}/{tender_a}/exports", json={}).status_code == 422


# --- Excel content ---------------------------------------------------------------------------


def test_workbook_matches_the_estimate_screens() -> None:
    actor, tender_id, est = _approved("xc1@example.com")
    made = _create(actor, tender_id, estimate_id=est["id"]).json()
    wb = _book(_download(actor, made["id"]))
    api_est = actor.client.get(f"{API}/estimates/{est['id']}").json()
    api_lines = _lines(actor, est)
    summary = _summary(wb)
    totals = api_est["totals"]
    assert _dec(summary["Subtotal cost"]) == _dec(totals["subtotal_cost"])
    assert _dec(summary["Subtotal ex VAT"]) == _dec(totals["subtotal_ex_vat"])
    assert _dec(summary["VAT"]) == _dec(totals["vat_total"])
    assert _dec(summary["Total incl VAT"]) == _dec(totals["total_inc_vat"])
    assert summary["Status"] == "Approved"
    rows = _data_rows(wb["Lines"])
    assert len(rows) == len(api_lines) > 0
    priced = 0
    for row, line in zip(rows, api_lines, strict=True):
        assert row[0] == line["position"] and row[3] == line["description"]
        if line["status"] == "priced":
            priced += 1
            assert _dec(row[18]) == _dec(line["total"])
            assert _dec(row[15]) == _dec(line["cost"])
            assert _dec(row[17]) == _dec(line["markup"])
            assert row[6] == line["supplier_name"]
        else:
            assert row[18] is None
    assert priced >= 1
    first_sheet = [wb["Summary"]["A1"].value, wb["Lines"]["A1"].value]
    assert not any("DRAFT" in str(v) for v in first_sheet)
    assert "DRAFT" not in str(wb["Summary"].oddHeader.center.text or "")
    assert "-DRAFT" not in made["file_name"]


def test_draft_and_reviewed_workbooks_carry_every_draft_marking() -> None:
    actor, tender_id, est = _ready("xc2@example.com")
    for state in ("draft", "reviewed"):
        if state == "reviewed":
            assert _move(actor, est, "review").status_code == 200
        made = _create(actor, tender_id).json()
        assert made["label"] == state and "-DRAFT-" in made["file_name"]
        wb = _book(_download(actor, made["id"]))
        for name in ("Summary", "Lines"):
            ws = wb[name]
            assert str(ws["A1"].value).startswith("DRAFT")
            assert "DRAFT" in ws.oddHeader.center.text
        assert wb.properties.title.startswith("DRAFT")


def test_changed_source_shows_source_changed_without_amounts() -> None:
    actor, tender_id, est = _ready("xc3@example.com")
    priced = next(x for x in _lines(actor, est) if x["status"] == "priced")
    _set_quantity(priced, "123")
    made = _create(actor, tender_id).json()
    rows = _data_rows(_book(_download(actor, made["id"]))["Lines"])
    row = next(r for r in rows if r[0] == priced["position"])
    assert row[19] == "Source changed"
    assert all(row[i] is None for i in range(11, 19))


# --- frozen exports --------------------------------------------------------------------------

_VARIES = {"Generated at (UTC)", "Export reference"}


def test_approved_export_is_frozen_against_later_edits() -> None:
    from app.modules.suppliers.models import QuotationItem

    actor, tender_id, est = _approved("xf1@example.com")
    a = _create(actor, tender_id).json()
    a_bytes = _download(actor, a["id"]).content
    line = next(x for x in _lines(actor, est) if x["status"] == "priced")
    _set_quantity(line, "999")
    with session_scope() as s:
        for q in s.scalars(select(QuotationItem)):
            if q.unit_price is not None:
                q.unit_price = q.unit_price + Decimal("50")
    b = _create(actor, tender_id).json()
    wa, wb_ = _book(_download(actor, a["id"])), _book(_download(actor, b["id"]))
    assert _sheet_values(wa["Lines"]) == _sheet_values(wb_["Lines"])
    sa, sb = _summary(wa), _summary(wb_)
    assert {k: v for k, v in sa.items() if k not in _VARIES} == {
        k: v for k, v in sb.items() if k not in _VARIES
    }
    assert sa["Export reference"] != sb["Export reference"]
    # A newer draft, exported on its own, leaves the first file alone.
    draft = actor.client.post(f"{TENDERS}/{tender_id}/estimates").json()
    d = _create(actor, tender_id, estimate_id=draft["id"]).json()
    assert d["label"] == "draft"
    assert _download(actor, a["id"]).content == a_bytes


def test_export_waits_for_the_lock_and_sees_a_consistent_reopened_state() -> None:
    actor, tender_id, est = _ready("xf2@example.com")
    assert _move(actor, est, "review").status_code == 200
    line = next(x for x in _lines(actor, est) if x["status"] == "priced")
    ctx = OrgContext(UUID(actor.me["user_id"]), UUID(actor.org_id), Role.ADMIN)
    holding = threading.Event()
    release = threading.Event()

    def hold_and_change() -> None:
        from app.api import estimate_inputs

        with session_scope() as s:
            suppliers.lock_tender(s, ctx, UUID(tender_id))
            holding.set()
            live = estimate_inputs.live_inputs(s, ctx, UUID(tender_id))
            estimates.reopen_estimate(s, ctx, UUID(est["id"]), live)
            estimates.price_line(
                s, ctx, UUID(line["id"]), estimates.LinePut(manual_unit_price="77.5"), live
            )
            release.wait(30)

    with ThreadPoolExecutor(2) as pool:
        holder = pool.submit(hold_and_change)
        assert holding.wait(30)
        export = pool.submit(lambda: _create(actor, tender_id))
        # Best effort: the assertions after the release are what prove the ordering, this only
        # gives the export request time to reach the lock.
        time.sleep(0.5)
        assert not export.done(), "the export did not wait for the tender lock"
        release.set()
        holder.result(30)
        made = export.result(30)
    assert made.status_code == 201, made.text
    assert made.json()["label"] == "draft"
    wb = _book(_download(actor, made.json()["id"]))
    api_est = actor.client.get(f"{API}/estimates/{est['id']}").json()
    assert _summary(wb)["Status"] == "Draft"
    assert _dec(_summary(wb)["Total incl VAT"]) == _dec(api_est["totals"]["total_inc_vat"])
    api_lines = _lines(actor, est)
    rows = _data_rows(wb["Lines"])
    for row, ln in zip(rows, api_lines, strict=True):
        if ln["status"] == "priced":
            assert _dec(row[18]) == _dec(ln["total"])
        else:
            assert row[18] is None


def test_export_waits_for_the_lock_and_sees_an_estimate_approved_meanwhile() -> None:
    actor, tender_id, est = _ready("xf3@example.com")
    assert _move(actor, est, "review").status_code == 200
    ctx = OrgContext(UUID(actor.me["user_id"]), UUID(actor.org_id), Role.ADMIN)
    holding = threading.Event()
    release = threading.Event()

    def hold_and_approve() -> None:
        from app.api import estimate_inputs

        with session_scope() as s:
            suppliers.lock_tender(s, ctx, UUID(tender_id))
            holding.set()
            live = estimate_inputs.live_inputs(s, ctx, UUID(tender_id))
            estimates.approve_estimate(s, ctx, UUID(est["id"]), live)
            release.wait(30)

    with ThreadPoolExecutor(2) as pool:
        holder = pool.submit(hold_and_approve)
        assert holding.wait(30)
        export = pool.submit(lambda: _create(actor, tender_id))
        time.sleep(0.5)  # best effort; the label below proves the ordering
        assert not export.done(), "the export did not wait for the tender lock"
        release.set()
        holder.result(30)
        made = export.result(30)
    assert made.status_code == 201, made.text
    assert made.json()["label"] == "approved"
    assert "DRAFT" not in made.json()["file_name"]
    assert _summary(_book(_download(actor, made.json()["id"])))["Status"] == "Approved"


# --- checklist PDF ---------------------------------------------------------------------------


def _seed_requirements(actor, upload_body: dict, rows: list[dict]) -> None:
    from app.modules.checklist import public as checklist
    from app.modules.pipeline.models import ExtractionRun

    ctx = OrgContext(UUID(actor.me["user_id"]), UUID(actor.org_id), Role.ADMIN)
    with session_scope() as s:
        run = ExtractionRun(
            organisation_id=ctx.organisation_id,
            document_version_id=UUID(upload_body["version"]["id"]),
            job_id=UUID(str(upload_body["job_id"])),
            kind="requirements",
        )
        s.add(run)
        s.flush()
        checklist.replace_unreviewed(s, ctx, run.document_version_id, run.id, rows)


def _req(description: str, source_order: int, rtype: str = "document", **extra) -> dict:
    return {
        "description": description, "requirement_type": rtype, "mandatory_status": "mandatory",
        "evidence_quote": f"quote {description}", "source_order": source_order, **extra,
    }  # fmt: skip


def _checklist_tender(email: str):
    actor = register(email)
    tender_id = make_tender(actor)
    first = upload(actor, tender_id, make_pdf(["one"]), "first.pdf").json()
    second = upload(actor, tender_id, make_pdf(["two"]), "second.pdf").json()
    return actor, tender_id, first, second


def _pdf_text(response) -> str:
    assert response.status_code == 200, response.text
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        # the big DRAFT watermark would interleave its letters with the text; leave it out
        return "\n".join(
            p.filter(lambda o: o.get("object_type") != "char" or o["size"] < 30).extract_text()
            or ""
            for p in pdf.pages
        )


def _all_requirements(actor, tender_id: str) -> dict[str, dict]:
    rows = actor.client.get(f"{TENDERS}/{tender_id}/requirements", params={"limit": 500}).json()
    return {r["description"]: r for r in rows}


def test_checklist_pdf_orders_rows_and_handles_review_states() -> None:
    actor, tender_id, first, second = _checklist_tender("xp1@example.com")
    _seed_requirements(actor, first, [_req("FIRSTLATE", 5), _req("FIRSTEARLY", 3)])
    _seed_requirements(
        actor,
        second,
        [
            _req("SECONDREJECT", 4),
            _req("SECONDCONFIRM", 0),
            _req("SECONDCORRECT", 2, "eligibility"),
            _req("SECONDPENDING", 3),
        ],
    )
    by = _all_requirements(actor, tender_id)
    base = f"{API}/requirements"
    assert actor.client.post(f"{base}/{by['FIRSTEARLY']['id']}/confirm").status_code == 200
    assert actor.client.post(f"{base}/{by['FIRSTLATE']['id']}/confirm").status_code == 200
    assert actor.client.post(f"{base}/{by['SECONDCONFIRM']['id']}/confirm").status_code == 200
    assert actor.client.post(f"{base}/{by['SECONDREJECT']['id']}/reject").status_code == 200
    r = actor.client.patch(f"{base}/{by['SECONDCORRECT']['id']}", json={"description": "FIXED"})
    assert r.status_code == 200, r.text
    made = _create(actor, tender_id, kind="checklist_pdf")
    assert made.status_code == 201, made.text
    assert made.json()["label"] == "draft" and made.json()["row_count"] == 5
    text = _pdf_text(_download(actor, made.json()["id"]))
    assert "SECONDREJECT" not in text and "1 rejected, not listed" in text
    assert "SECONDPENDING" in text and "Not yet reviewed (1)" in text
    assert text.index("Not yet reviewed (1)") < text.index("SECONDPENDING")
    assert "FIXED" in text and "Eligibility" in text
    # document order, then source order inside a document. SECONDCONFIRM has the lowest source
    # order of all, so ordering by source order alone would put it first.
    listed = actor.client.get(f"{TENDERS}/{tender_id}/documents").json()
    names = [d["current_version"]["filename"] for d in listed]
    assert sorted(names) == ["first.pdf", "second.pdf"]
    first_block = ["FIRSTEARLY", "FIRSTLATE"]
    second_block = ["SECONDCONFIRM"]
    blocks = {"first.pdf": first_block, "second.pdf": second_block}
    expected = [t for n in names for t in blocks[n]]
    assert [text.index(t) for t in expected] == sorted(text.index(t) for t in expected)
    assert text.index("FIRSTEARLY") < text.index("FIRSTLATE")  # source order inside a document
    assert "first.pdf" in text and "second.pdf" in text


def test_checklist_ignores_superseded_document_versions() -> None:
    actor, tender_id, first, _ = _checklist_tender("xp2@example.com")
    _seed_requirements(actor, first, [_req("OLDVERSIONROW", 0)])
    again = actor.client.post(
        f"{API}/documents/{first['document']['id']}/versions",
        files={"file": ("first.pdf", make_pdf(["new"]), "application/pdf")},
    )
    assert again.status_code == 201, again.text
    made = _create(actor, tender_id, kind="checklist_pdf").json()
    assert made["row_count"] == 0
    assert "OLDVERSIONROW" not in _pdf_text(_download(actor, made["id"]))


# --- download headers, filename, missing file ------------------------------------------------


def test_download_headers_and_safe_filename() -> None:
    actor = register("xh1@example.com")
    tender_id = make_tender(actor, 'x"; evil=1 /../é')
    pdf = _create(actor, tender_id, kind="checklist_pdf").json()
    r = _download(actor, pdf["id"])
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    disposition = r.headers["content-disposition"]
    assert disposition == f'attachment; filename="{pdf["file_name"]}"'
    assert all(c.isalnum() or c in '-.=" ;' for c in disposition), disposition
    assert (
        '"' not in pdf["file_name"] and "/" not in pdf["file_name"] and ";" not in pdf["file_name"]
    )
    assert pdf["file_name"].isascii() and pdf["file_name"].endswith(".pdf")
    assert r.headers["x-content-type-options"] == "nosniff"
    assert r.headers["cache-control"] == "private, no-store"
    assert r.headers["content-security-policy"] == "sandbox"
    assert int(r.headers["content-length"]) == len(r.content) == pdf["file_size"]

    actor2, tender2, est = _approved("xh2@example.com")
    xl = _create(actor2, tender2).json()
    r2 = _download(actor2, xl["id"])
    assert r2.headers["content-type"] == XLSX
    assert r2.headers["content-disposition"] == f'attachment; filename="{xl["file_name"]}"'
    assert int(r2.headers["content-length"]) == len(r2.content)


def test_missing_stored_file_is_404_and_other_organisations_are_refused() -> None:
    actor = register("xh3@example.com")
    tender_id = make_tender(actor)
    made = _create(actor, tender_id, kind="checklist_pdf").json()
    with session_scope() as s:
        key = s.scalar(select(Export.file_key).where(Export.id == UUID(made["id"])))
    get_storage().delete(key)
    assert _download(actor, made["id"]).status_code == 404
    other = register("xh3b@example.com")
    assert other.client.get(f"{EXPORTS}/{made['id']}").status_code == 404
    assert _download(other, made["id"]).status_code == 404


def test_archived_tender_can_still_be_exported() -> None:
    actor, tender_id, _ = _approved("xa1@example.com")
    assert actor.client.post(f"{TENDERS}/{tender_id}/archive").status_code == 200
    assert _create(actor, tender_id).status_code == 201
    assert _create(actor, tender_id, kind="checklist_pdf").status_code == 201


# --- audit, rate limit, caps -----------------------------------------------------------------


def test_audit_rows_carry_only_allowlisted_keys() -> None:
    actor, tender_id, est = _approved("xu1@example.com")
    made = _create(actor, tender_id).json()
    assert _download(actor, made["id"]).status_code == 200
    with session_scope() as s:
        rows = s.scalars(
            select(AuditEvent).where(
                AuditEvent.action.in_(["export_created", "export_downloaded"]),
                AuditEvent.organisation_id == UUID(actor.org_id),
            )
        ).all()
    by = {r.action: r.event_metadata for r in rows}
    assert set(by) == {"export_created", "export_downloaded"}
    assert set(by["export_created"]) <= {"kind", "estimate_id", "version", "status", "count"}
    assert by["export_created"]["estimate_id"] == est["id"]
    assert set(by["export_downloaded"]) <= {"kind"}
    assert made["file_name"] not in str(by)


def test_eleventh_create_in_a_minute_is_429() -> None:
    actor = register("xl1@example.com")
    tender_id = make_tender(actor)
    codes = [_create(actor, tender_id, kind="checklist_pdf").status_code for _ in range(11)]
    assert codes[:10] == [201] * 10 and codes[10] == 429


def test_too_many_rows_is_422(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.modules.reporting.service.MAX_ROWS", 3)
    actor, tender_id, _ = _priced_setup("xl2@example.com")
    r = _create(actor, tender_id)
    assert r.status_code == 422 and r.json()["error"]["code"] == "export_too_large"
    # the checklist cap counts non-rejected rows only
    other, tender2, up, _ = _checklist_tender("xl3@example.com")
    _seed_requirements(other, up, [_req(f"ROW{i}", i) for i in range(4)])
    big = _create(other, tender2, kind="checklist_pdf")
    assert big.status_code == 422 and big.json()["error"]["code"] == "export_too_large"


# --- injection -------------------------------------------------------------------------------


def test_formula_and_markup_text_stays_inert() -> None:
    from app.modules.estimates.models import EstimateLine
    from app.modules.estimating.models import BoqItem
    from app.modules.suppliers.models import Supplier

    actor, tender_id, est = _priced_setup("xi1@example.com")
    with session_scope() as s:
        for item in (*s.scalars(select(BoqItem)), *s.scalars(select(EstimateLine))):
            item.description = '=HYPERLINK("http://example.com","x")'
        for n, sup in enumerate(s.scalars(select(Supplier).order_by(Supplier.name_key))):
            sup.name = f"@SUM({n})"
    _resolve_all(actor, est)
    assert _move(actor, est, "review").status_code == 200
    assert _move(actor, est, "approve").status_code == 200
    made = _create(actor, tender_id).json()
    wb = _book(_download(actor, made["id"]))
    rows = _data_rows(wb["Lines"])
    assert all(str(r[3]).startswith("'=HYPERLINK") for r in rows)
    suppliers_seen = {r[6] for r in rows if r[6]}
    assert suppliers_seen and all(str(s).startswith("'@SUM") for s in suppliers_seen)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate)
                assert c.hyperlink is None

    actor2, tender2, first, _ = _checklist_tender("xi2@example.com")
    _seed_requirements(actor2, first, [_req('<img src="file:///etc/passwd"/>', 0)])
    pdf = _create(actor2, tender2, kind="checklist_pdf").json()
    assert 'src="file:///etc/passwd"' in _pdf_text(_download(actor2, pdf["id"]))
