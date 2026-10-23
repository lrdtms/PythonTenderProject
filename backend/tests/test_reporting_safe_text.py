"""The text-safety funnel for exports: hostile strings must be inert in Excel and PDF."""

import io
import json
import re
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph

from app.modules.reporting import words
from app.modules.reporting.safe_text import (
    export_filename,
    pdf_plain,
    pdf_text,
    xlsx_header_text,
    xlsx_text,
)

CORPUS = json.loads(
    (Path(__file__).parent / "fixtures" / "exports" / "injection_corpus.json").read_text()
)
ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")
STYLE = getSampleStyleSheet()["Normal"]


def _cleaned(s: str) -> str:
    return ILLEGAL.sub("", s)


@pytest.mark.parametrize("s", CORPUS, ids=range(len(CORPUS)))
def test_xlsx_text_is_inert_and_accepted_by_openpyxl(s):
    out = xlsx_text(s)
    assert not ILLEGAL.search(out)
    assert len(out) <= 32_000
    cleaned = _cleaned(s)
    cut = cleaned if len(cleaned) <= 32_000 - 4 else cleaned[: 32_000 - 4] + "..."
    risky = cut.lstrip().startswith(("=", "+", "-", "@")) or cut.startswith(("\t", "\r"))
    assert out.startswith("'") == risky or (cut.startswith("'") and not risky)
    if risky:
        assert out == "'" + cut
    else:
        assert out == cut
    wb = Workbook()
    wb.active["A1"] = out
    buf = io.BytesIO()
    wb.save(buf)
    cell = load_workbook(io.BytesIO(buf.getvalue())).active["A1"]
    assert cell.data_type == "s" or out == ""


def test_xlsx_text_none_and_limit():
    assert xlsx_text(None) == ""
    assert len(xlsx_text("=" + "x" * 40_000)) <= 32_000
    assert xlsx_text("x" * 50, limit=20) == "x" * 16 + "..."


@pytest.mark.parametrize("s", CORPUS, ids=range(len(CORPUS)))
def test_pdf_text_is_markup_safe(s):
    out = pdf_text(s)
    assert len(out) <= 2_000
    assert "<" not in out.replace("&lt;", "").replace("&gt;", "")
    assert ">" not in out.replace("&gt;", "")
    assert "<img" not in out
    Paragraph(out, STYLE)
    for ch in out:
        ch.encode("cp1252")


def test_pdf_text_specifics():
    assert pdf_text("≤ 5 Ω") == "? 5 ?"
    assert "?" in pdf_text("ok \U0001f600")
    assert pdf_text("Ngiyabonga, môre, ê, ü") == "Ngiyabonga, môre, ê, ü"
    assert pdf_text("a & b") == "a &amp; b"
    long = pdf_text("r" * 5_000)
    assert long.endswith("... (truncated, see source)")
    assert len(long) <= 2_000


def test_header_text():
    assert xlsx_header_text("A & B") == "A && B"
    assert xlsx_header_text("&Z&F") == "&&Z&&F"
    assert len(xlsx_header_text("T" * 300)) == 60
    for s in CORPUS:
        out = xlsx_header_text(s)
        assert not ILLEGAL.search(out)
        assert not re.search(r"(?<!&)&(?!&)", out.replace("&&", ""))


def test_pdf_plain():
    assert pdf_plain("Smith & Sons") == "Smith & Sons"
    assert pdf_plain("≤") == "?"
    assert pdf_plain("a\x00b") == "ab"
    assert len(pdf_plain("x" * 1000)) <= 300


TITLES = [
    "Greenvale Hall / Phase 2",
    'x"; rm',
    "../../etc",
    "",
    "t" * 200,
    "Ngiyabonga môre ≤",
    "...",
]


@pytest.mark.parametrize("title", TITLES)
@pytest.mark.parametrize("draft", [True, False])
def test_export_filename(title, draft):
    xlsx = export_filename(title, "estimate", 3, draft, date(2026, 10, 23))
    pdf = export_filename(title, "checklist", None, draft, date(2026, 10, 23))
    pattern = r"^[A-Za-z0-9][A-Za-z0-9.\-]*$"
    assert re.match(pattern, xlsx) and xlsx.endswith("-20261023.xlsx")
    assert re.match(pattern, pdf) and pdf.endswith("-20261023.pdf")
    assert ("-DRAFT-" in xlsx) is draft and ("-DRAFT-" in pdf) is draft
    assert "-estimate-v3" in xlsx and "-checklist" in pdf
    slug = xlsx.split("-estimate-")[0]
    assert 1 <= len(slug) <= 40


def test_export_filename_empty_title_is_tender():
    assert export_filename("", "checklist", None, False, date(2026, 1, 2)).startswith("tender-")


FLAG_CODES = [
    "currency_differs", "currency_missing", "vat_basis_unknown", "price_missing",
    "quantity_missing", "unit_missing", "unit_mismatch", "unreviewed_source", "source_changed",
]  # fmt: skip


@pytest.mark.parametrize("code", FLAG_CODES)
def test_flag_text_known(code):
    text = words.flag_text(code)
    assert text and text != words.flag_text("nope")


def test_flag_text_fallback():
    assert words.flag_text("nope") == "Needs attention"


def _section(src: str, marker: str) -> str:
    return src.split(marker)[1].split("\ndef ")[0]


def test_every_flag_the_estimates_module_emits_has_a_word():
    src = (
        Path(__file__).parent.parent / "app" / "modules" / "estimates" / "service.py"
    ).read_text()
    found = set(re.findall(r'flags\["(\w+)"\]\s*=\s*True', src))
    found |= set(re.findall(r'out\.append\("(\w+)"\)', _section(src, "def _warnings")))
    found |= set(re.findall(r'flags\.append\("(\w+)"\)', _section(src, "def _line_dto")))
    assert set(FLAG_CODES) <= found
    fallback = words.flag_text("nope")
    missing = [c for c in found if words.flag_text(c) == fallback]
    assert not missing


def test_other_words():
    assert words.vat_basis_text(True) != words.vat_basis_text(False)
    assert words.vat_basis_text(None)
    for s in ("draft", "reviewed", "approved"):
        assert words.estimate_status_text(s)
    for s in ("unreviewed", "confirmed", "corrected", "rejected"):
        assert words.review_status_text(s)
    for s in ("mandatory", "conditional", "informational"):
        assert words.mandatory_text(s)
    for s in ("supplier", "manual", "excluded", "unresolved"):
        assert words.price_source_text(s)
    for s in ("cost", "selling_price"):
        assert words.markup_basis_text(s)


def test_header_truncation_keeps_ampersands_paired():
    for s in ("x" * 59 + "&", "&" * 61, "&" * 300):
        out = xlsx_header_text(s)
        assert len(out) <= 120 and out.replace("&&", "").count("&") == 0


def test_pdf_text_truncation_edges():
    out = pdf_text("<" * 3000)
    assert out.endswith("... (truncated, see source)") and out.startswith("&lt;")
    assert len(out) <= 2_000
    assert len(pdf_text("x" * 50, limit=10)) <= 10


def test_export_filename_validation_and_case():
    d = date(2026, 1, 2)
    with pytest.raises(ValueError):
        export_filename("t", "pptx", None, False, d)
    with pytest.raises(ValueError):
        export_filename("t", "estimate", None, False, d)
    assert export_filename("Greenvale Hall", "checklist", None, False, d).startswith(
        "greenvale-hall-"
    )


def test_xlsx_text_unicode_space_before_starter():
    assert xlsx_text("\u00a0=1+1").startswith("'")
