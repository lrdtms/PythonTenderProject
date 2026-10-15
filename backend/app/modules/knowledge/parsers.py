"""Turn files into pages of text and tables. No database access."""

from __future__ import annotations

import csv
import io
import math
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from typing import Any

import pdfplumber
import pymupdf
from docx import Document as DocxDocument
from openpyxl import load_workbook

from app.config import get_settings
from app.kernel import ProcessingFailed

from .ocr import OcrFailed, OcrRunner

MIN_EMBEDDED_CHARS = 20  # fewer characters than this and the page is treated as scanned
LOW_CONFIDENCE = 60.0
MIN_OCR_DPI = 100  # below this Tesseract output is not worth having
MAX_SHEET_ROWS = 100_000
MAX_TABLES_PER_PAGE = 30

Progress = Callable[[int, int], None]


@dataclass
class ParsedPage:
    page_number: int  # physical 1-based page index (PDF) or 1-based sheet index
    text: str
    label: str | None = None  # sheet name
    text_source: str = "embedded"  # embedded | ocr | none
    ocr_confidence: float | None = None
    status: str = "ok"  # ok | empty | ocr_failed | ocr_unavailable | low_confidence
    tables: list[dict[str, Any]] = field(default_factory=list)


def _cell(value: object) -> str:
    return "" if value is None else str(value).strip()


# --- PDF -------------------------------------------------------------------------------------


def parse_pdf(
    data: bytes, ocr: OcrRunner | None, ocr_dpi: int = 300, progress: Progress | None = None
) -> list[ParsedPage]:
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001 - any parser error means an unreadable file
        raise ProcessingFailed("FILE_CORRUPT", "The PDF could not be opened.") from exc
    with doc:
        if doc.needs_pass or doc.is_encrypted:
            raise ProcessingFailed("PDF_ENCRYPTED", "The PDF is password protected.")
        total = doc.page_count
        if total == 0:
            raise ProcessingFailed("FILE_CORRUPT", "The PDF has no pages.")
        try:
            plumber = pdfplumber.open(io.BytesIO(data))
        except Exception:  # noqa: BLE001 - tables are best effort
            plumber = None
        pages: list[ParsedPage] = []
        try:
            for index, page in enumerate(doc):
                pages.append(_parse_pdf_page(page, index, plumber, ocr, ocr_dpi))
                if progress:
                    progress(index + 1, total)
        finally:
            if plumber is not None:
                plumber.close()
        return pages


def _parse_pdf_page(
    page: pymupdf.Page, index: int, plumber: Any, ocr: OcrRunner | None, ocr_dpi: int
) -> ParsedPage:
    text = page.get_text("text")
    number = index + 1
    if len(text.strip()) >= MIN_EMBEDDED_CHARS:
        return ParsedPage(number, text, tables=_pdf_tables(plumber, index))
    # Scanned (or blank) page: OCR if we can.
    if not page.get_images() and not text.strip():
        return ParsedPage(number, "", text_source="none", status="empty")
    if ocr is None or not ocr.available():
        return ParsedPage(number, text, text_source="none", status="ocr_unavailable")
    dpi = _render_dpi(page.rect.width, page.rect.height, ocr_dpi)
    if dpi is None:  # too large to render safely, even at the lowest useful DPI
        return ParsedPage(number, "", text_source="none", status="ocr_failed")
    try:
        png = page.get_pixmap(dpi=dpi).tobytes("png")
        ocr_text, confidence = ocr.run(png)
    except OcrFailed:
        return ParsedPage(number, "", text_source="none", status="ocr_failed")
    status = "low_confidence" if confidence < LOW_CONFIDENCE else "ok"
    return ParsedPage(number, ocr_text, text_source="ocr", ocr_confidence=confidence, status=status)


def _render_dpi(width_pt: float, height_pt: float, dpi: int) -> int | None:
    """The DPI to render a page for OCR so the image stays within the pixel cap, or None when even
    MIN_OCR_DPI would exceed it (a forged or poster-sized page must not allocate gigabytes)."""
    max_pixels = get_settings().ocr_max_megapixels * 1_000_000
    pixels = (width_pt / 72 * dpi) * (height_pt / 72 * dpi)
    if pixels <= max_pixels:
        return dpi
    lowered = math.floor(dpi * math.sqrt(max_pixels / pixels))
    return lowered if lowered >= MIN_OCR_DPI else None


def _pdf_tables(plumber: Any, index: int) -> list[dict[str, Any]]:
    if plumber is None:
        return []
    try:
        found = plumber.pages[index].extract_tables()
    except Exception:  # noqa: BLE001 - a page whose tables cannot be read still has its text
        return []
    return [
        {"kind": "table", "rows": [[_cell(c) for c in row] for row in table]}
        for table in found[:MAX_TABLES_PER_PAGE]
        if table
    ]


# --- spreadsheets, CSV, DOCX -----------------------------------------------------------------


class _CellBudget:
    """Caps the columns per row and the cells per file (settings), so a small file cannot make
    the worker build billions of cells. Exceeding a cap is a permanent SHEET_TOO_LARGE."""

    def __init__(self) -> None:
        settings = get_settings()
        self.max_columns = settings.xlsx_max_columns
        self.remaining = settings.xlsx_max_cells

    def take(self, row: Iterable[object]) -> list[str]:
        cells = [_cell(c) for c in row]
        while cells and not cells[-1]:  # trailing empty cells carry nothing
            cells.pop()
        if len(cells) > self.max_columns:
            raise ProcessingFailed("SHEET_TOO_LARGE", "A worksheet has too many columns.")
        self.remaining -= len(cells)
        if self.remaining < 0:
            raise ProcessingFailed("SHEET_TOO_LARGE", "The file has too many cells.")
        return cells


def parse_xlsx(data: bytes, progress: Progress | None = None) -> list[ParsedPage]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ProcessingFailed("FILE_CORRUPT", "The workbook could not be opened.") from exc
    pages: list[ParsedPage] = []
    sheets = wb.worksheets
    budget = _CellBudget()
    for number, ws in enumerate(sheets, start=1):
        # The declared <dimension> is untrusted: read-only mode pads every row up to it, so a
        # forged A1:XFD3 turns a tiny sheet into 16,384 cells per row.
        ws.reset_dimensions()
        rows: list[list[str]] = []
        for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row_index > MAX_SHEET_ROWS:
                raise ProcessingFailed("SHEET_TOO_LARGE", "A worksheet has too many rows.")
            rows.append(budget.take(row))
        text = "\n".join(" | ".join(c for c in r if c) for r in rows if any(r))
        pages.append(
            ParsedPage(
                number, text, label=ws.title, status="ok" if text else "empty",
                tables=[{"kind": "sheet", "first_row": 1, "rows": rows}] if rows else [],
            )
        )  # fmt: skip
        if progress:
            progress(number, len(sheets))
    wb.close()
    return pages


def parse_csv(data: bytes, progress: Progress | None = None) -> list[ParsedPage]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    budget = _CellBudget()
    rows: list[list[str]] = []
    for row in csv.reader(io.StringIO(text), dialect):
        if len(rows) >= MAX_SHEET_ROWS:
            raise ProcessingFailed("SHEET_TOO_LARGE", "The CSV has too many rows.")
        rows.append(budget.take(row))
    body = "\n".join(" | ".join(c for c in r if c) for r in rows if any(r))
    if progress:
        progress(1, 1)
    return [
        ParsedPage(
            1, body, label="csv", status="ok" if body else "empty",
            tables=[{"kind": "sheet", "first_row": 1, "rows": rows}] if rows else [],
        )
    ]  # fmt: skip


def parse_docx(data: bytes, progress: Progress | None = None) -> list[ParsedPage]:
    """Word files have no fixed pages; the whole document is treated as page 1."""
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ProcessingFailed("FILE_CORRUPT", "The document could not be opened.") from exc
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = [
        {"kind": "table", "rows": [[_cell(c.text) for c in row.cells] for row in t.rows]}
        for t in doc.tables
    ]
    text = "\n".join(paragraphs)
    if progress:
        progress(1, 1)
    return [
        ParsedPage(
            1, text, label="docx (no page numbers)", status="ok" if text else "empty", tables=tables
        )
    ]


def parse_document(
    file_format: str,
    data: bytes,
    ocr: OcrRunner | None = None,
    ocr_dpi: int = 300,
    progress: Progress | None = None,
) -> list[ParsedPage]:
    if file_format == "pdf":
        return parse_pdf(data, ocr, ocr_dpi, progress)
    if file_format == "xlsx":
        return parse_xlsx(data, progress)
    if file_format == "csv":
        return parse_csv(data, progress)
    if file_format == "docx":
        return parse_docx(data, progress)
    raise ProcessingFailed("UNSUPPORTED_FORMAT", f"Cannot process {file_format} files.")
