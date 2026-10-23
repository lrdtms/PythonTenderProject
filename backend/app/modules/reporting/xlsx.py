"""Excel renderer for an estimate version: a pure function of `EstimateExportInput`.

Every string reaches a cell through `safe_text.xlsx_text`; headers and footers through
`xlsx_header_text`. No formulas, no hyperlinks. A `Decimal` is a numeric cell only when it has at
most 15 significant digits (openpyxl stores a float), otherwise it is an inline text cell.
"""

import io
from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.reporting import words
from app.modules.reporting.safe_text import xlsx_header_text, xlsx_text
from app.modules.reporting.schemas import EstimateExportInput, ExportLine

SUMMARY_SHEET = "Summary"
LINES_SHEET = "Lines"
LINE_COLUMNS: tuple[str, ...] = (
    "Line",
    "Section",
    "Item no",
    "Description",
    "Unit",
    "Quantity",
    "Supplier",
    "Price source",
    "Quoted unit price",
    "Quote currency",
    "Quote VAT basis",
    "Unit price ex VAT",
    "Labour rate",
    "Material cost",
    "Labour cost",
    "Cost",
    "Markup (%)",
    "Markup",
    "Line total (ex VAT)",
    "Status",
    "Notes",
)

AMOUNT = "#,##0.00"
RATE = "#,##0.0000"
PERCENT = "0.00"
MAX_DIGITS = 15
RED = "C00000"
GREY = "D9D9D9"
_LINE_WIDTHS = (
    6, 16, 9, 44, 8, 11, 22, 16, 14, 10, 14, 14, 12, 14, 12, 14, 10, 12, 16, 15, 48,
)  # fmt: skip

_PRICE_SOURCE = {
    "supplier": "Supplier quote",
    "manual": "Manual price",
    "excluded": "Excluded",
    "unresolved": "Not priced",
}
_LINE_STATUS = {
    "priced": "Priced",
    "excluded": "Excluded",
    "unresolved": "Not priced",
    "changed": "Source changed",
}
_MARKUP_BASIS = {
    "cost": "Percentage on cost",
    "selling_price": "Percentage on selling price",
}
_STATUS = {"draft": "Draft", "reviewed": "Reviewed: not yet approved", "approved": "Approved"}
_NOTE = (
    "Amounts are values, not formulas. Quotes that include VAT are converted to ex VAT on the "
    "stated rate; nothing is converted between currencies."
)
_LONG_NOTE = "Some values have more than 15 digits and are shown as text."


def is_draft(status: str) -> bool:
    return status != "approved"


def _utc(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=UTC)
    return value.astimezone(UTC)


def _when(value: datetime) -> str:
    return _utc(value).strftime("%Y-%m-%d %H:%M UTC")


class _Book:
    """A workbook plus the one write helper every cell goes through."""

    def __init__(self) -> None:
        self.wb = Workbook()
        self.long_values = False

    def put(self, ws: Worksheet, row: int, col: int, value: object, fmt: str | None = None) -> None:
        if value is None:
            return
        cell = ws.cell(row=row, column=col)
        if isinstance(value, datetime):
            value = _when(value)
        if isinstance(value, bool):
            value = "Yes" if value else "No"
        if isinstance(value, Decimal) and (
            not value.is_finite() or len(value.normalize().as_tuple().digits) > MAX_DIGITS
        ):
            self.long_values = True
            value = format(value, "f") if value.is_finite() else str(value)
        if isinstance(value, str):
            cell.value = xlsx_text(value)
            cell.data_type = "s"
        else:
            cell.value = value
            if fmt:
                cell.number_format = fmt


def _banner(ws: Worksheet, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    cell = ws["A1"]
    cell.fill = PatternFill("solid", fgColor=RED)
    cell.font = Font(bold=True, color="FFFFFF", size=12)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _page_setup(ws: Worksheet, data: EstimateExportInput, landscape: bool) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.oddFooter.left.text = f"{xlsx_header_text(data.tender_title)}, version {data.version}"
    ws.oddFooter.right.text = "Page &P of &N"
    if is_draft(data.status):
        ws.oddHeader.center.text = f"DRAFT: version {data.version} is not approved"
        ws.oddHeader.center.font = "Calibri,Bold"
        ws.oddHeader.center.size = 14
        ws.oddHeader.center.color = RED
        ws.sheet_properties.tabColor = RED


def _person(name: str | None, when: datetime | None, missing: str) -> str:
    if not name:
        return missing
    return f"{name}, {_when(when)}" if when else name


def _notes(line: ExportLine) -> str | None:
    if line.status == "excluded":
        return line.exclude_reason
    return "; ".join(words.flag_text(f) for f in line.flags) or None


def _line_row(
    book: _Book, ws: Worksheet, row: int, line: ExportLine, default_markup: Decimal
) -> None:
    markup_percent = line.markup_percent if line.markup_percent is not None else default_markup
    put = book.put
    put(ws, row, 1, line.position, "General")
    put(ws, row, 2, line.section)
    put(ws, row, 3, line.item_no)
    put(ws, row, 4, line.description)
    put(ws, row, 5, line.unit)
    put(ws, row, 6, line.quantity, RATE)
    put(ws, row, 7, line.supplier_name)
    put(ws, row, 8, _PRICE_SOURCE[line.price_source])
    if line.status != "excluded":
        put(ws, row, 9, line.quoted_unit_price, RATE)
        put(ws, row, 10, line.quoted_currency)
        if line.quoted_unit_price is not None:
            basis = "Not stated"
            if line.quoted_includes_vat is not None:
                basis = words.vat_basis_text(line.quoted_includes_vat)
            put(ws, row, 11, basis)
    if line.status == "priced":
        put(ws, row, 12, line.unit_price_ex_vat, RATE)
        put(ws, row, 13, line.labour_rate, RATE)
        put(ws, row, 14, line.material_cost, AMOUNT)
        put(ws, row, 15, line.labour_cost, AMOUNT)
        put(ws, row, 16, line.cost, AMOUNT)
        put(ws, row, 17, markup_percent, PERCENT)
        put(ws, row, 18, line.markup, AMOUNT)
        put(ws, row, 19, line.total, AMOUNT)
    put(ws, row, 20, _LINE_STATUS[line.status])
    put(ws, row, 21, _notes(line))
    if line.status == "excluded":
        grey = PatternFill("solid", fgColor=GREY)
        for col in range(1, len(LINE_COLUMNS) + 1):
            ws.cell(row=row, column=col).fill = grey


def _write_lines(book: _Book, ws: Worksheet, data: EstimateExportInput) -> None:
    columns = len(LINE_COLUMNS)
    header = 1
    if is_draft(data.status):
        book.put(ws, 1, 1, f"DRAFT: version {data.version} is not approved. Do not submit.")
        _banner(ws, columns)
        header = 2
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    for i, name in enumerate(LINE_COLUMNS, start=1):
        book.put(ws, header, i, name)
        cell = ws.cell(row=header, column=i)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    last = header
    for line in data.lines:
        last += 1
        _line_row(book, ws, last, line, data.markup_percent)
    for i, width in enumerate(_LINE_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = f"A{header + 1}"
    ws.auto_filter.ref = f"A{header}:{ws.cell(row=last, column=columns).coordinate}"
    ws.print_title_rows = f"{header}:{header}"
    _page_setup(ws, data, landscape=True)


def _write_summary(book: _Book, ws: Worksheet, data: EstimateExportInput) -> None:
    put = book.put
    row = 1
    if is_draft(data.status):
        put(ws, 1, 1, f"DRAFT: version {data.version} is not approved. Do not submit.")
        _banner(ws, 2)
        row = 2
    put(ws, row, 1, "Tender estimate")
    ws.cell(row=row, column=1).font = Font(bold=True, size=14)
    row += 1

    def pair(label: str, value: object, fmt: str | None = None) -> int:
        nonlocal row
        put(ws, row, 1, label)
        ws.cell(row=row, column=1).font = Font(bold=True)
        put(ws, row, 2, value, fmt)
        row += 1
        return row - 1

    pair("Tender", data.tender_title)
    pair("Reference", data.tender_reference)
    pair("Client", data.client_name or "Not stated")
    pair("Estimate version", data.version, "General")
    pair("Status", _STATUS[data.status])
    pair("Approved by", _person(data.approved_by_name, data.approved_at, "Not approved"))
    pair("Reviewed by", _person(data.reviewed_by_name, data.reviewed_at, "Not reviewed"))
    pair("Currency", data.currency)
    pair("VAT rate (%)", data.vat_rate, PERCENT)
    pair("Markup basis", _MARKUP_BASIS[data.markup_basis])
    pair("Default markup (%)", data.markup_percent, PERCENT)
    pair("Lines priced", data.priced, "General")
    pair("Lines excluded", data.excluded, "General")
    pair("Lines not priced", data.unresolved, "General")
    pair("Lines changed", data.changed, "General")
    totals = data.totals
    labels = (
        ("Subtotal cost", "subtotal_cost"),
        ("Markup", "markup_total"),
        ("Subtotal ex VAT", "subtotal_ex_vat"),
        ("VAT", "vat_total"),
        ("Total incl VAT", "total_inc_vat"),
    )
    total_row = row
    for label, field in labels:
        total_row = pair(
            label, "Not available" if totals is None else getattr(totals, field), AMOUNT
        )
    if data.incomplete:
        skipped = data.unresolved + data.changed
        put(
            ws,
            total_row,
            3,
            f"Incomplete: excludes {skipped} {'line' if skipped == 1 else 'lines'}",
        )
        ws.cell(row=total_row, column=3).font = Font(bold=True, color=RED)
    ws.cell(row=total_row, column=1).font = Font(bold=True)
    ws.cell(row=total_row, column=2).font = Font(bold=True)
    pair("Generated by", data.generated_by_name)
    pair("Generated at (UTC)", data.generated_at)
    pair("Export reference", str(data.export_id))
    put(ws, row, 1, _NOTE)
    row += 1
    if book.long_values:
        put(ws, row, 1, _LONG_NOTE)
    for letter, width in (("A", 30), ("B", 46), ("C", 32)):
        ws.column_dimensions[letter].width = width
    _page_setup(ws, data, landscape=False)


def render_estimate_xlsx(data: EstimateExportInput) -> bytes:
    if data.generated_at is None or data.export_id is None:
        raise ValueError("generated_at and export_id must be set before rendering")
    book = _Book()
    summary = book.wb.active
    if summary is None:
        raise ValueError("new workbook has no active sheet")
    summary.title = SUMMARY_SHEET
    lines = book.wb.create_sheet(LINES_SHEET)
    _write_lines(book, lines, data)  # first, so the Summary knows about long values
    _write_summary(book, summary, data)
    props = book.wb.properties
    prefix = "DRAFT: " if is_draft(data.status) else ""
    props.title = xlsx_text(
        f"{prefix}Tender estimate: {data.tender_title}, version {data.version}", 255
    )
    props.creator = "IntelTPS"
    stamp = _utc(data.generated_at).replace(tzinfo=None)
    props.created = stamp
    props.modified = stamp
    out = io.BytesIO()
    book.wb.save(out)
    return out.getvalue()
